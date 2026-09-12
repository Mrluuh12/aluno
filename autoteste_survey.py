#!/usr/bin/env python3
"""
Autoverificação do gerador de relatórios do MeshMapper.

Roda a ferramenta inteira sobre a captura de exemplo e confere as saídas.
Serve para responder "a instalação está boa?" sem precisar de um arquivo
de campo:

    python autoteste_survey.py

A suíte completa (505 testes) fica no projeto principal; aqui o objetivo
é outro — provar que ESTA cópia, nesta máquina, gera o que promete.
"""
import sys, io, zipfile, tempfile, shutil
from pathlib import Path

AQUI = Path(__file__).resolve().parent
EXEMPLO = AQUI / "exemplos" / "meshmapper_exemplo.json"
# Captura feita PARADA (o MeshMapper ligado numa repetidora). Produto
# diferente: censo de vizinhos, não rastro de calor.
REPETIDORA = AQUI / "exemplos" / "meshmapper_repetidora.json"
falhas = []


def checa(condicao, descricao, detalhe=""):
    print(("  ok    " if condicao else "  FALHOU ") + descricao
          + (f"  [{detalhe}]" if detalhe and not condicao else ""))
    if not condicao:
        falhas.append(descricao)


def main():
    print("Autoverificação — gerador de relatórios do MeshMapper\n")

    print("1. Dependências")
    for mod, pacote in (("matplotlib", "matplotlib"), ("numpy", "numpy"),
                        ("scipy", "scipy"), ("pptx", "python-pptx"),
                        ("openpyxl", "openpyxl"), ("lxml", "lxml")):
        try:
            __import__(mod); checa(True, pacote)
        except ImportError as e:
            checa(False, pacote, str(e))
    try:
        import tkinter; checa(True, "tkinter (a janela)")
    except ImportError:
        # Não é impeditivo: a linha de comando funciona sem.
        print("  aviso  tkinter ausente — a janela não abre, "
              "mas a linha de comando funciona")

    print("\n2. Arquivos")
    checa((AQUI / "rajant_monitor.py").exists(), "rajant_monitor.py (o motor)")
    checa((AQUI / "survey_meshmapper.py").exists(), "survey_meshmapper.py")
    checa(EXEMPLO.exists(), "exemplos/meshmapper_exemplo.json")
    if not REPETIDORA.exists():
        print("  aviso  exemplos/meshmapper_repetidora.json ausente")
    marca = AQUI / "marca" / "anglo_azul.png"
    if marca.exists():
        checa(True, "marca/ (logos)")
    else:
        print("  aviso  marca/ ausente — o deck sai sem logo")
    if falhas:
        print(f"\n{len(falhas)} problema(s). Pare aqui.")
        return 1

    sys.path.insert(0, str(AQUI))
    import survey_meshmapper as sm
    import rajant_monitor as rm

    print("\n3. Leitura da captura de exemplo")
    sv, am, pr = rm.ler_meshmapper(str(EXEMPLO))
    checa(len(am) > 0, f"{len(am)} amostras lidas")
    checa(len(pr) > 0, f"{len(pr)} leituras de vizinho")
    checa(sv.get("movel") is not None, f"equipamento: {sv.get('movel')}")
    # Signal é dBm (negativo e grande); rssi do arquivo é SNR (positivo).
    com_sinal = [a for a in am if a["sinal"] is not None]
    checa(all(a["sinal"] < -20 for a in com_sinal), "RSSI em dBm")
    checa(all(0 <= a["snr"] < 70 for a in com_sinal if a["snr"] is not None),
          "SNR em dB")
    checa(all(a["ruido"] == a["sinal"] - a["snr"] for a in com_sinal
              if a["snr"] is not None), "ruído recuperado de signal − snr")
    # 0 dBm não existe num rádio de malha: é o "ainda não medi" do
    # arquivo. Deixar passar fazia o vizinho zerado ganhar a eleição de
    # cobertura e pintar o ponto de verde máximo.
    checa(all(p["sinal"] != 0 for p in pr), "nenhum vizinho com 0 dBm")
    checa(not rm.captura_parada(am), "reconhecida como captura em trajeto")

    print("\n4. Leitura de uma captura feita parada")
    if REPETIDORA.exists():
        sv2, am2, pr2 = rm.ler_meshmapper(str(REPETIDORA))
        checa(rm.captura_parada(am2),
              f"{sv2.get('movel')} reconhecida como parada "
              f"({rm.extensao_da_captura(am2)['raio_m']:g} m de raio)")
        censo = rm.censo_vizinhos(pr2, len(am2))
        r = rm.resumo_do_censo(censo)
        checa(len(censo) > 1, f"censo com {len(censo)} vizinhos "
                              f"({r['infra']} de infraestrutura)")
        checa(all(c["presenca"] is not None for c in censo),
              "presença calculada para todos")
        dados, nome = rm.gerar_kml_pontos_fixos(
            [{"nome": sv2["movel"], **{k: rm.extensao_da_captura(am2)[k]
                                       for k in ("lat", "lon", "raio_m")},
              "pontos": len(am2), "censo": censo, "resumo": r,
              "inicio": sv2["inicio"], "fim": sv2["fim"]}])
        from xml.dom.minidom import parseString
        doc = zipfile.ZipFile(io.BytesIO(dados)).read("doc.kml").decode()
        try:
            parseString(doc); checa(True, "KMZ da vizinhança é XML válido")
        except Exception as e:
            checa(False, "KMZ da vizinhança é XML válido", str(e))
    else:
        print("  aviso  exemplos/meshmapper_repetidora.json ausente — "
              "pulei a verificação da captura parada")

    print("\n5. Geração")
    tmp = Path(tempfile.mkdtemp())
    try:
        feitos = sm.gerar([str(EXEMPLO)], str(tmp), aviso=lambda t: None)
        checa(len(feitos) >= 1, f"{len(feitos)} arquivo(s) gerado(s)")
        checa(all(f.exists() for f in feitos), "todos existem no disco")

        kmz = [f for f in feitos if f.suffix == ".kmz"]
        if kmz:
            z = zipfile.ZipFile(kmz[0])
            doc = z.read("doc.kml").decode()
            checa("<GroundOverlay>" in doc, "KMZ tem o rastro de calor")
            checa("<LatLonBox>" in doc, "o calor está georreferenciado")
            pngs = [n for n in z.namelist() if n.endswith(".png")]
            checa(bool(pngs), f"raster embutido ({len(pngs)})")
            if pngs:
                import matplotlib; matplotlib.use("Agg")
                import matplotlib.image as mpimg
                img = mpimg.imread(io.BytesIO(z.read(pngs[0])))
                transp = (img[..., 3] < 0.02).mean()
                checa(transp > 0.3,
                      f"o calor só cobre onde passou ({transp*100:.0f}% "
                      f"transparente)")

        ppt = [f for f in feitos if f.suffix == ".pptx"]
        if ppt:
            from pptx import Presentation
            p = Presentation(str(ppt[0]))
            n = len(list(p.slides))
            checa(n >= 4, f"PPT com {n} slides")

        xls = [f for f in feitos if f.suffix == ".xlsx"]
        if xls:
            from openpyxl import load_workbook
            wb = load_workbook(str(xls[0]))
            checa("Vizinhos" in wb.sheetnames,
                  f"Excel com as abas {wb.sheetnames}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print()
    if falhas:
        print(f"{len(falhas)} problema(s):")
        for f in falhas: print(f"  - {f}")
        return 1
    print("Tudo certo. A instalação está boa.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
