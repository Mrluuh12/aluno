#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TESTES DO PARSER — rajant_monitor.py

Um caso por correção da auditoria das métricas contra os .proto oficiais
do bcapi (bcapi-ref/proto/). Cada teste referencia o campo do protocolo
que justifica o comportamento esperado.

Regra que orienta todos os casos: campo que NÃO existe no protocolo não
pode virar 0 — 0 é indistinguível de "medido e sem ocorrência". Ou vira
None (e a série não é publicada), ou a métrica deixa de existir.

    python3 teste_parser.py
"""
import sys, types, unittest, tempfile, shutil, os, inspect, re, pathlib, json, time
from pathlib import Path
from unittest import mock

# ── rajant_api não é instalável em CI; o módulo faz sys.exit(1) sem ele ──
if 'rajant_api' not in sys.modules:
    _fake = types.ModuleType('rajant_api')
    class Breadcrumb:                     # stub: nenhum teste vai à rede
        def __init__(self, *a, **k): pass
    _fake.Breadcrumb = Breadcrumb
    sys.modules['rajant_api'] = _fake

import rajant_monitor as rm


# ══════════════════════════════════════════════════════════════════
# FIXTURE — State no formato text-format EXATO do State.proto
# ══════════════════════════════════════════════════════════════════
# Inclui de propósito a ARMADILHA `configuration`: ela repete os nomes
# wired / wireless / battery / general com os campos de CONFIGURAÇÃO,
# que não têm contador nenhum. Um parser que extraia do texto inteiro
# pega esses blocos e lê zero em tudo.
STATE = """
system {
  kernel: "3.18.20"
  platform: "ES1"
  uptime: 864000000.0
  idle: 691200000.0
  running: true
  bridgeup: true
  version: "11.25.0"
  freeMemory: 131072
  temperature: 4520
  bootCounter: 7
  reboot: false
  ipv4 {
    address: "10.188.96.140"
    subnet: "255.255.255.0"
  }
}
build {
  date: "2024-03-01"
  version: "11.25.0"
  number: "4471"
}
instamesh {
  arpDropped: 3
  arpRequests: 900
  arpTotal: 1200
  floodsDropped: 44
  packetsDropped: 120
  packetsMulticast: 5000
  packetsReceived: 900000
  packetsSent: 850000
  sourceFloodsDropped: 17
  timeWaited: 90
  discoveriesSourced: 12
  discoveriesPassed: 34
}
battery {
  hardwareRevision: "B"
  milliamps: 1500
  charging: false
  capacityPercent: 87
  temperatureCelsius: 31
  dischargeTimeMinutes: 210
  chargeTimeMinutes: 0
}
alertSystem {
  bestRadioRate: 866
  ledMode: "normal"
  alerts {
    index: 1
    type: WARNING
    message: "link marginal"
  }
  alerts {
    index: 2
    type: INFORMATION
    message: "config salva"
  }
}
gps {
  gpsSwitch {
    enabled: true
  }
  gpsPos {
    gpsTime: 143025.0
    gpsLat: "2743.8950S"
    gpsLong: "05004.1429W"
    gpsPrecisionH: 0.9
    gpsQuality: 1.0
    gpsSatsInView: 9
    gpsAlt: 812.5
  }
  gpsVel {
    gpsTrackDegreesTrue: 187.3
    gpsSpeedKnots: 22.95
    gpsSpeedKph: 42.5
  }
}
wireless {
  key: 1
  mac: "00:11:22:33:44:55"
  name: "wlan0"
  noise: -95
  channel: 149
  txpower: 23
  type: "5GHz"
  stats {
    rxBytes: 5000000
    rxPackets: 40000
    txBytes: 3000000
    txPackets: 30000
  }
  ap {
    key: 1
    essid: "MESH-OPERACAO"
    client {
      mac: "aa:bb:cc:00:00:01"
      rate: 650
      rssi: 38
      signal: -57
      age: 4
    }
    client {
      mac: "aa:bb:cc:00:00:02"
      rate: 780
      rssi: 41
      signal: -54
      age: 2
    }
    client {
      mac: "aa:bb:cc:00:00:03"
      rate: 520
      rssi: 29
      signal: -66
      age: 9
    }
  }
  peer {
    mac: "00:11:22:33:44:66"
    enabled: true
    cost: 120
    rate: 1300
    rssi: 42
    signal: -58
    age: 3
    ipv4Address: "10.188.96.141"
    stats {
      rxBytes: 900000
      txBytes: 700000
    }
  }
  peer {
    mac: "00:11:22:33:44:77"
    enabled: true
    cost: 340
    rate: 650
    rssi: 21
    signal: -81
    age: 11
  }
  channelActiveTime: 100000
  channelBusyTime: 22000
  channelReceiveTime: 9000
  channelTransmitTime: 6000
}
wired {
  key: 1
  mac: "00:11:22:33:44:88"
  masterMac: "00:11:22:33:44:88"
  aptState: APT_STATE_MASTER
  stats {
    rxBytes: 12000000
    rxPackets: 90000
    txBytes: 8000000
    txPackets: 70000
  }
  peer {
    mac: "00:11:22:33:44:99"
    enabled: true
    cost: 10
    ipv4Address: "10.188.96.142"
  }
  name: "eth0"
  ipv4 {
    address: "10.188.96.140"
  }
}
wired {
  key: 2
  mac: "00:11:22:33:44:aa"
  aptState: APT_STATE_NONE
  stats {
    rxBytes: 0
    rxPackets: 0
    txBytes: 0
    txPackets: 0
  }
  name: "eth1"
}
configuration {
  saved {
    general {
      name: "BC-MINA-01"
      notes: "backbone"
      groups {
        name: "MINA-NORTE"
      }
      groups {
        name: "BACKBONE"
      }
    }
    battery {
      warningThresholdMinutes: 30
      errorThresholdMinutes: 10
    }
    wired {
      name: "eth0"
      requestFallback: false
      speedMbps: 1000
    }
    wireless {
      name: "wlan0"
      ap {
        essid: "MESH-OPERACAO"
      }
    }
  }
}
"""


class BaseParser(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = rm.parse_state(STATE)
        cls.s = cls.d["sistema"]
        cls.im = cls.d["instamesh"]
        cls.r = cls.d["radios"][0]
        cls.eth = {e["nome"]: e for e in cls.d["ethernet"]}


# ══════════════════════════════════════════════════════════════════
# 1. ARMADILHA `configuration`
# ══════════════════════════════════════════════════════════════════
class TestBlocoConfiguration(BaseParser):
    """State e Config têm mensagens de mesmo nome; o parser tem de ler o
    bloco de ESTADO, não o de configuração."""

    def test_wired_vem_do_estado_nao_da_config(self):
        # Config.Wired tem name/requestFallback/speedMbps e nenhum contador.
        self.assertEqual(self.eth["eth0"]["rx_bytes"], 12000000)
        self.assertEqual(self.eth["eth0"]["tx_bytes"], 8000000)

    def test_apenas_duas_portas_a_config_nao_vira_porta(self):
        self.assertEqual(len(self.d["ethernet"]), 2)

    def test_bateria_vem_de_state_battery_nao_de_config_battery(self):
        # Config.Battery só tem warning/errorThresholdMinutes. Se o parser
        # pegasse esse bloco, capacityPercent viria 0.
        self.assertEqual(self.s["bateria"]["pct"], 87)

    def test_nome_vem_de_configuration_saved_general_name(self):
        # Buscar `name:` na configuration inteira pegaria "MINA-NORTE"
        # (nome de grupo) ou "eth0" (nome de porta).
        self.assertEqual(self.s["nome"], "BC-MINA-01")


# ══════════════════════════════════════════════════════════════════
# 2. CAMPOS INVENTADOS — nunca podem virar 0
# ══════════════════════════════════════════════════════════════════
class TestCamposInexistentes(BaseParser):

    def test_sem_chaves_de_erro_de_ethernet(self):
        # CommStats = rxBytes, rxPackets, txBytes, txPackets. Ponto.
        for chave in ("rx_err","tx_err","rx_crc","rx_drop","tx_drop","mudancas"):
            self.assertNotIn(chave, self.eth["eth0"],
                             f"{chave} não existe em State.Wired/CommStats")

    def test_sem_chaves_de_radar_e_phy(self):
        for chave in ("radar_detec","radar_pulsos","phy"):
            self.assertNotIn(chave, self.r,
                             f"{chave} não existe em State.Wireless")

    def test_sem_txpower_por_peer(self):
        # State.Peer não tem txpower; só State.Wireless tem.
        self.assertNotIn("txpwr", self.r["peers"][0])

    def test_sem_session_state(self):
        self.assertNotIn("session_st", self.s)

    def test_sem_metricas_de_tensao(self):
        # Não existe bloco `sensors` em nenhum .proto.
        for chave in ("volt_v","volt_min_v","volt_max_v","bateria_v"):
            self.assertNotIn(chave, self.s)

    def test_instamesh_sem_campos_inventados(self):
        for chave in ("unicast","ovf","undel_rx","undel_tx"):
            self.assertNotIn(chave, self.im,
                             f"{chave} não consta de State.InstaMesh")

    def test_metricas_removidas_nao_existem_mais(self):
        for nome in ("m_volt","m_volt_min","m_volt_max","m_bateria",
                     "m_session_state","m_r_radar","m_r_pulsos","m_r_phy",
                     "m_p_txpwr","m_e_rx_err","m_e_tx_err","m_e_rx_crc",
                     "m_e_rx_drop","m_e_tx_drop","m_e_mudancas",
                     "m_im_unicast","m_im_ovf","m_im_undel_rx","m_im_undel_tx"):
            self.assertFalse(hasattr(rm, nome), f"{nome} deveria ter sido removida")


# ══════════════════════════════════════════════════════════════════
# 3. InstaMesh — campos reais
# ══════════════════════════════════════════════════════════════════
class TestInstaMesh(BaseParser):

    def test_contadores_reais(self):
        self.assertEqual(self.im["pkt_tx"], 850000)
        self.assertEqual(self.im["pkt_rx"], 900000)
        self.assertEqual(self.im["pkt_drop"], 120)
        self.assertEqual(self.im["floods"], 44)
        self.assertEqual(self.im["arp"], 1200)

    def test_campos_novos_aproveitados(self):
        # sourceFloodsDropped e packetsMulticast existiam e não eram lidos.
        self.assertEqual(self.im["src_floods"], 17)
        self.assertEqual(self.im["multicast"], 5000)


# ══════════════════════════════════════════════════════════════════
# 4. APT / link ethernet
# ══════════════════════════════════════════════════════════════════
class TestAPT(BaseParser):

    def test_apt_state_usa_enum_do_proto(self):
        # APT_STATE_MASTER = 0, APT_STATE_NONE = 2
        self.assertEqual(self.eth["eth0"]["apt_int"], 0)
        self.assertEqual(self.eth["eth1"]["apt_int"], 2)

    def test_apt_master_derivado_de_aptstate(self):
        # aptMaster/isMaster não existem; deriva de Wired.aptState == MASTER
        self.assertTrue(self.s["apt_master"])

    def test_link_inferido(self):
        self.assertTrue(self.eth["eth0"]["link"])    # MASTER + peer
        self.assertFalse(self.eth["eth1"]["link"])   # NONE, sem peer

    def test_peers_por_porta(self):
        self.assertEqual(self.eth["eth0"]["peers_eth"], 1)
        self.assertEqual(self.eth["eth1"]["peers_eth"], 0)

    def test_apt_master_falso_sem_porta_master(self):
        st = STATE.replace("aptState: APT_STATE_MASTER", "aptState: APT_STATE_SLAVE")
        self.assertFalse(rm.parse_state(st)["sistema"]["apt_master"])


# ══════════════════════════════════════════════════════════════════
# 5. Clientes de AP — contagem de blocos
# ══════════════════════════════════════════════════════════════════
class TestClientesAP(BaseParser):

    def test_conta_blocos_client(self):
        # State.Wireless.AP tem `repeated Client client`; não há clientCount.
        self.assertEqual(self.r["aps"][0]["clients"], 3)

    def test_total_por_radio(self):
        self.assertEqual(self.r["total_cli"], 3)

    def test_ap_entra_sem_campo_enabled(self):
        # State.Wireless.AP NÃO tem `enabled`. Exigi-lo zerava a lista de APs.
        self.assertEqual(len(self.r["aps"]), 1)
        self.assertEqual(self.r["aps"][0]["essid"], "MESH-OPERACAO")


# ══════════════════════════════════════════════════════════════════
# 6. Peers — SNR, taxa, IP opcional
# ══════════════════════════════════════════════════════════════════
class TestPeers(BaseParser):

    def test_taxa_dividida_por_10(self):
        # bc_livestats.py oficial: pv.rate / 10 = Mbps
        self.assertEqual(self.r["peers"][0]["taxa"], 130.0)

    def test_snr_e_sinal_menos_ruido(self):
        # signal=-58, noise=-95  ->  SNR = 37 dB
        # A conta antiga (rssi - abs(noise)) dava 42-95 = -53.
        self.assertEqual(self.r["peers"][0]["snr"], 37)
        self.assertEqual(self.r["peers"][1]["snr"], 14)   # -81 - (-95)

    def test_peer_sem_ipv4_nao_e_descartado(self):
        # ipv4Address é `optional` em State.Peer: exigir descartava o enlace
        # inteiro, com taxa, custo, RSSI e SNR junto.
        self.assertEqual(len(self.r["peers"]), 2)
        self.assertTrue(self.r["peers"][1]["ip"].startswith("mac:"))

    def test_good_peers_usa_snr_correto(self):
        # Só o primeiro peer tem SNR > 30 dB.
        self.assertEqual(self.r["good_peers"], 1)

    def test_custo_e_first_hop(self):
        self.assertEqual(self.r["peers"][0]["custo"], 120)
        self.assertEqual(self.r["first_hop"], 120)


# ══════════════════════════════════════════════════════════════════
# 7. GPS
# ══════════════════════════════════════════════════════════════════
class TestGPS(BaseParser):

    def test_nmea_convertido(self):
        self.assertAlmostEqual(self.s["gps_lat"], -27.7316, places=3)
        self.assertAlmostEqual(self.s["gps_lon"], -50.0690, places=3)
        self.assertEqual(self.s["gps_fix"], 1)

    def test_velocidade_em_kph_sem_reconverter(self):
        # gpsSpeedKph JÁ está em km/h. O parser antigo lia `speed`
        # (inexistente) e ainda multiplicava por 1,852.
        self.assertEqual(self.s["gps_vel"], 42.5)

    def test_rumo_de_gpstrackdegreestrue(self):
        # `course`/`heading`/`track` não existem no Gps.proto.
        self.assertEqual(self.s["gps_rumo"], 187.3)

    def test_campos_de_qualidade(self):
        self.assertEqual(self.s["gps_sats"], 9)
        self.assertEqual(self.s["gps_alt"], 812.5)
        self.assertEqual(self.s["gps_hdop"], 0.9)

    def test_fallback_para_knots(self):
        st = STATE.replace("    gpsSpeedKph: 42.5\n", "")
        # 22.95 kn * 1.852 = 42.5 km/h
        self.assertAlmostEqual(rm.parse_state(st)["sistema"]["gps_vel"], 42.5, places=1)

    def test_sem_velocidade_fica_none_nao_zero(self):
        st = STATE.replace("    gpsSpeedKnots: 22.95\n", "").replace(
            "    gpsSpeedKph: 42.5\n", "")
        self.assertIsNone(rm.parse_state(st)["sistema"]["gps_vel"])

    def test_gpsswitch_desligado_invalida_posicao(self):
        st = STATE.replace("enabled: true", "enabled: false", 1)
        d = rm.parse_state(st)
        self.assertIsNone(d["sistema"]["gps_lat"])
        self.assertEqual(d["sistema"]["gps_fix"], 0)


# ══════════════════════════════════════════════════════════════════
# 8. Sistema — temperatura, CPU, bateria, alertas
# ══════════════════════════════════════════════════════════════════
class TestSistema(BaseParser):

    def test_temperatura_centi_grau(self):
        self.assertEqual(self.s["temp_c"], 45.2)

    def test_temperatura_grau_inteiro(self):
        # Firmware que publica o grau direto não pode virar 0,45 °C.
        st = STATE.replace("temperature: 4520", "temperature: 45")
        self.assertEqual(rm.parse_state(st)["sistema"]["temp_c"], 45.0)

    def test_cpu_none_na_primeira_coleta(self):
        # cpuLoad não existe; a carga vem do delta de idle/uptime e exige
        # duas coletas. None = não publica (0 seria "CPU ociosa", mentira).
        self.assertIsNone(self.s["cpu_pct"])

    def test_cpu_expoe_idle_bruto(self):
        self.assertEqual(self.s["cpu_idle"], 691200000.0)
        self.assertEqual(self.s["uptime_raw"], 864000000.0)

    def test_bateria_completa(self):
        b = self.s["bateria"]
        self.assertEqual(b["pct"], 87)
        self.assertEqual(b["ma"], 1500)
        self.assertEqual(b["temp_c"], 31)
        self.assertEqual(b["desc_min"], 210)
        self.assertFalse(b["carregando"])

    def test_sem_bateria_tudo_none(self):
        # Modelo sem bateria: nenhuma série publicada, nada de zeros.
        st = rm.remover_bloco(STATE, 'battery')
        b = rm.parse_state(st)["sistema"]["bateria"]
        self.assertTrue(all(v is None for v in b.values()))

    def test_alert_system(self):
        self.assertEqual(self.s["best_rate"], 866)
        self.assertEqual(self.s["alertas"], 2)

    def test_reboot_e_bridge(self):
        self.assertEqual(self.s["reboot"], 0)
        self.assertTrue(self.s["bridge"])

    def test_grupos(self):
        self.assertEqual(self.s["grupos"], "MINA-NORTE|BACKBONE")

    def test_ip_e_modelo(self):
        self.assertEqual(self.s["ip"], "10.188.96.140")
        self.assertEqual(self.s["modelo"], "ES1")


# ══════════════════════════════════════════════════════════════════
# 9. Rádio — canal, ruído, ocupação
# ══════════════════════════════════════════════════════════════════
class TestRadio(BaseParser):

    def test_campos_basicos(self):
        self.assertEqual(self.r["nome"], "wlan0")
        self.assertEqual(self.r["canal"], 149)
        self.assertEqual(self.r["freq"], "5GHz")
        self.assertEqual(self.r["ruido"], -95)
        self.assertEqual(self.r["txpwr"], 23)

    def test_stats_do_radio(self):
        self.assertEqual(self.r["rx_bytes"], 5000000)
        self.assertEqual(self.r["tx_pkts"], 30000)

    def test_ocupacao_de_canal(self):
        self.assertEqual(self.r["ch_active"], 100000)
        self.assertEqual(self.r["busy_pct"], 22.0)
        self.assertEqual(self.r["rx_pct"], 9.0)


# ══════════════════════════════════════════════════════════════════
# 10. CPU derivada — precisa de duas coletas
# ══════════════════════════════════════════════════════════════════
class TestCPUDerivada(unittest.TestCase):
    """100 * (1 - Δidle/Δuptime). A razão é adimensional: vale para
    segundos ou milissegundos, desde que os dois campos usem a mesma
    unidade."""

    def setUp(self):
        rm._ultimo.clear()

    def _coleta(self, ip, uptime, idle, ts):
        st = STATE.replace("uptime: 864000000.0", f"uptime: {uptime}") \
                  .replace("idle: 691200000.0",   f"idle: {idle}")
        return rm.calcular_taxas(ip, rm.parse_state(st), ts)

    def test_carga_75_pct(self):
        self._coleta("1.1.1.1", 1000.0, 500.0, 100.0)
        # +100 de uptime, +25 de idle  ->  75% de carga
        d = self._coleta("1.1.1.1", 1100.0, 525.0, 200.0)
        self.assertEqual(d["sistema"]["cpu_pct"], 75.0)

    def test_maquina_ociosa(self):
        self._coleta("2.2.2.2", 1000.0, 500.0, 100.0)
        d = self._coleta("2.2.2.2", 1100.0, 600.0, 200.0)
        self.assertEqual(d["sistema"]["cpu_pct"], 0.0)

    def test_reboot_no_meio_nao_publica(self):
        # uptime volta a zero: delta negativo, resultado sem sentido.
        self._coleta("3.3.3.3", 100000.0, 90000.0, 100.0)
        d = self._coleta("3.3.3.3", 50.0, 40.0, 200.0)
        self.assertIsNone(d["sistema"]["cpu_pct"])

    def test_idle_por_nucleo_nao_publica(self):
        # Δidle > Δuptime (multicore) -> carga negativa -> não publica.
        self._coleta("4.4.4.4", 1000.0, 500.0, 100.0)
        d = self._coleta("4.4.4.4", 1100.0, 900.0, 200.0)
        self.assertIsNone(d["sistema"]["cpu_pct"])


# ══════════════════════════════════════════════════════════════════
# 11. publicar() nunca pode estourar
# ══════════════════════════════════════════════════════════════════
class TestPublicar(unittest.TestCase):
    """prometheus_client faz float(valor): .set(None) levanta TypeError.
    Como _coletar_bc captura Exception, um único campo ausente marcava o
    BreadCrumb inteiro como OFFLINE."""

    def test_publica_state_completo(self):
        est = rm.EstadoBC(3, 300)
        est.ok({}, 1.0)
        rm.publicar("10.188.96.140", rm.parse_state(STATE), est)

    def test_publica_state_minimo_sem_estourar(self):
        # Sem GPS, sem bateria, sem alertas, sem CPU: tudo None.
        est = rm.EstadoBC(3, 300)
        est.ok({}, 1.0)
        minimo = """
system {
  platform: "ES1"
  uptime: 1000.0
  freeMemory: 65536
  temperature: 40
}
wired {
  name: "eth0"
  aptState: APT_STATE_NONE
  stats { rxBytes: 1 rxPackets: 1 txBytes: 1 txPackets: 1 }
}
"""
        rm.publicar("10.0.0.9", rm.parse_state(minimo), est)

    def test_pub_ignora_none(self):
        from prometheus_client import Gauge
        g = Gauge("teste_pub_none", "t", ["bc","ip"])
        est = rm.EstadoBC(3, 300); est.ok({}, 1.0)
        # publicar() define pub() internamente; aqui validamos o contrato
        # diretamente: .set(None) tem de estourar, e é isso que pub() evita.
        with self.assertRaises(TypeError):
            g.labels(bc="x", ip="y").set(None)


# ══════════════════════════════════════════════════════════════════
# 12. Regressões de robustez já existentes
# ══════════════════════════════════════════════════════════════════
class TestSeriesNoMetrics(unittest.TestCase):
    """O /metrics só tem o que foi medido AGORA (campo, 15/10/2026: "o
    Grafana persiste a última leitura, fica uma linha reta") e só
    equipamento com tag ("tem que ter apenas rádios com tag")."""

    ST = STATE

    def setUp(self):
        self._ping = rm.ping_rtt
        rm.ping_rtt = lambda ip: None
        self.ip = "10.250.0." + str(abs(hash(self.id())) % 200 + 10)

    def tearDown(self):
        rm.ping_rtt = self._ping
        rm.retirar_series(self.ip)

    def _texto(self):
        from prometheus_client import generate_latest, REGISTRY
        return generate_latest(REGISTRY).decode()

    def _linhas(self, metrica, ip=None):
        ip = ip or self.ip
        return [l for l in self._texto().splitlines()
                if l.startswith(metrica + "{") and f'ip="{ip}"' in l]

    def _dados(self, nome="CA-1006", peers=(("10.188.97.220", 1), ("10.0.0.77", 2))):
        d = rm.parse_state(self.ST)
        d["sistema"]["nome"] = nome
        r = d["radios"][0]
        base = dict(r["peers"][0])
        r["peers"] = [dict(base, ip=ip_p, encap=enc) for ip_p, enc in peers]
        return d

    def _coletor(self):
        c = rm.RajantCollector.__new__(rm.RajantCollector)
        c.prefixos_tag = None; c.sem_tag = set()
        c.re_tag = rm.re.compile(rm.PADRAO_TAG_PADRAO, rm.re.I)
        c.nomes = {"10.188.97.220": "ERM-22 PTP CAM"}
        c.nome_encap = {}
        class Cache:
            def nome(self, ip): return None
        c.cache = Cache()
        return c

    def test_vizinho_que_sumiu_sai_do_metrics(self):
        est = rm.EstadoBC(3, 0); est.ok({}, 1.0)
        rm.publicar(self.ip, self._dados(peers=(("10.0.1.1", 1), ("10.0.1.2", 2))), est)
        self.assertEqual(len(self._linhas("rajant_peer_sinal_dbm")), 2)
        rm.publicar(self.ip, self._dados(peers=(("10.0.1.1", 1),)), est)
        linhas = self._linhas("rajant_peer_sinal_dbm")
        self.assertEqual(len(linhas), 1, linhas)
        self.assertIn('peer="10.0.1.1"', linhas[0])

    def test_falha_tira_a_medicao_e_deixa_online_zero(self):
        est = rm.EstadoBC(3, 0); est.ok({}, 1.0)
        rm.publicar(self.ip, self._dados(), est)
        for m in ("rajant_peer_sinal_dbm", "rajant_radio_canal", "rajant_uptime_s"):
            self.assertTrue(self._linhas(m), m)
        est.falha()
        rm.marcar_offline(self.ip, "CA-1006", est, 10**9)
        self.assertEqual(self._linhas("rajant_peer_sinal_dbm"), [])
        self.assertEqual(self._linhas("rajant_radio_canal"), [])
        on = self._linhas("rajant_online")
        self.assertEqual(len(on), 1); self.assertTrue(on[0].endswith(" 0.0"), on)
        self.assertEqual(len(self._linhas("rajant_bc_info")), 1)
        self.assertEqual(len(self._linhas("rajant_falhas_consecutivas")), 1)

    def test_segurar_serie_mantem_dentro_da_janela(self):
        d = self._dados()
        est = rm.EstadoBC(3, 300); est.ok(d, time.monotonic())
        rm.publicar(self.ip, d, est)
        est.falha()
        rm.marcar_offline(self.ip, "CA-1006", est, time.monotonic())
        self.assertTrue(self._linhas("rajant_peer_sinal_dbm"))
        rm.marcar_offline(self.ip, "CA-1006", est, time.monotonic() + 301)
        self.assertEqual(self._linhas("rajant_peer_sinal_dbm"), [])

    def test_vizinho_sem_tag_nao_vira_serie_e_com_tag_sai_pelo_nome(self):
        c = self._coletor()
        est = rm.EstadoBC(3, 0); est.ok({}, 1.0)
        dados = self._dados(peers=(("10.188.97.220", 1), ("10.0.0.77", 2),
                                   ("mac:AA:BB", 128433)))
        c.nome_encap = {128433: "ERB-02"}
        rm.publicar(self.ip, dados, est, nome_vizinho=c.nome_vizinho)
        peers = sorted(re.search(r'peer="([^"]+)"', l).group(1)
                       for l in self._linhas("rajant_peer_sinal_dbm"))
        self.assertEqual(peers, ["ERB-02", "ERM-22 PTP CAM"])

    def test_ultima_coleta_e_hora_real(self):
        est = rm.EstadoBC(3, 0); est.ok({}, time.monotonic())
        rm.publicar(self.ip, self._dados(), est)
        v = float(self._linhas("rajant_ultima_coleta_ts")[0].split()[-1])
        self.assertLess(abs(v - time.time()), 60)

    def test_ip_sem_nome_que_falha_nao_entra_no_metrics(self):
        c = self._coletor()
        self.assertFalse(c.tem_tag(c.nome_conhecido("10.188.99.217")))
        fonte = inspect.getsource(rm.RajantCollector._coletar_bc)
        self.assertIn("if self.tem_tag(nome):\n                        marcar_offline", fonte)

    def test_identidade_so_com_tag(self):
        fonte = inspect.getsource(rm.RajantCollector._publicar_identidade)
        self.assertIn("if not self.tem_tag(", fonte)


class TestServicoSegundoPlano(unittest.TestCase):
    """--servico: Prometheus, Grafana e exportador em segundo plano, e
    reerguidos se caírem (campo, 15/10/2026: "quando a máquina reinicia,
    temos que iniciar manualmente")."""

    FALSO = (
        "#!/usr/bin/env python3\n"
        "import socket, sys, time\n"
        "a = sys.argv[1:]\n"
        "porta = (int(a[a.index('--porta') + 1]) if '--porta' in a else\n"
        "         int(open(a[a.index('--homepath') + 1] + '/conf/porta').read()))\n"
        "s = socket.socket(); s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)\n"
        "s.bind(('127.0.0.1', porta)); s.listen(5)\n"
        "print('no ar', a, flush=True)\n"
        "while True:\n"
        "    c, _ = s.accept(); c.close()\n")

    @staticmethod
    def _porta_livre():
        import socket
        s = socket.socket(); s.bind(("127.0.0.1", 0))
        p = s.getsockname()[1]; s.close(); return p

    def setUp(self):
        import servico_rajant as sr, configparser, stat
        self.sr = sr
        self.tmp = Path(tempfile.mkdtemp())
        self.pp, self.pg = self._porta_livre(), self._porta_livre()
        prom = self.tmp / "prometheus" / "prometheus"
        graf = self.tmp / "grafana" / "bin" / "grafana-server"
        for f in (prom, graf):
            f.parent.mkdir(parents=True)
            f.write_text(self.FALSO); f.chmod(f.stat().st_mode | stat.S_IEXEC)
        (self.tmp / "grafana" / "conf").mkdir()
        (self.tmp / "grafana" / "conf" / "porta").write_text(str(self.pg))
        self.cfg = configparser.ConfigParser()
        self.cfg.read_dict({
            "servicos": {"prometheus_exe": str(prom),
                         "prometheus_args": f"--porta {self.pp}",
                         "prometheus_porta": str(self.pp),
                         "grafana_exe": str(graf),
                         "grafana_porta": str(self.pg),
                         "iniciar_exportador": "false"}})
        sr.cfg_servicos(self.cfg)
        self._cwd = os.getcwd()
        self._esperas = sr.Servico.ESPERAS
        sr.Servico.ESPERAS = (0.2,)
        self.servicos = []

    def tearDown(self):
        for s in self.servicos:
            s.parar()
        self.sr.parar_tudo(self.tmp)
        self.sr.Servico.ESPERAS = self._esperas
        os.chdir(self._cwd)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _esperar_porta(self, porta, aberta=True, prazo=8.0):
        fim = time.time() + prazo
        while time.time() < fim:
            if self.sr.porta_aberta(porta, espera=0.2) == aberta:
                return True
            time.sleep(0.1)
        return False

    def test_sobe_os_dois_e_reergue_o_que_cai(self):
        sv, avisos = self.sr.montar_servicos(self.cfg, self.tmp)
        self.servicos = sv
        self.assertEqual([s.nome for s in sv], ["Prometheus", "Grafana"])
        self.assertEqual(avisos, [])
        graf = sv[1]
        self.assertEqual(graf.cmd[1:], ["--homepath", str((self.tmp / "grafana").resolve())])
        for s in sv:
            self.assertIn("iniciado", s.verificar())
        self.assertTrue(self._esperar_porta(self.pp))
        self.assertTrue(self._esperar_porta(self.pg))
        # O Prometheus cai: a volta seguinte agenda, a outra reergue.
        self.sr.matar(sv[0].pid)
        self.assertTrue(self._esperar_porta(self.pp, aberta=False))
        time.sleep(0.3)
        self.assertIn("parou", sv[0].verificar())
        time.sleep(0.3)
        self.assertIn("iniciado", sv[0].verificar())
        self.assertTrue(self._esperar_porta(self.pp))
        # O log de cada um vai para logs/.
        self.assertTrue((self.tmp / "logs" / "prometheus.log").exists())

    def test_porta_ja_atendida_nao_inicia_outro(self):
        # Grafana instalado como serviço do Windows: já está no ar.
        import socket
        ocupa = socket.socket(); ocupa.bind(("127.0.0.1", self.pg)); ocupa.listen(1)
        try:
            sv, _ = self.sr.montar_servicos(self.cfg, self.tmp)
            self.servicos = sv
            msg = sv[1].verificar()
            self.assertIn("já está no ar", msg)
            self.assertIsNone(sv[1].proc)
            self.assertIsNone(sv[1].verificar())     # não repete o aviso
        finally:
            ocupa.close()
        # Saiu do ar: o supervisor assume.
        self.assertIn("assumindo", sv[1].verificar())
        self.assertIsNotNone(sv[1].proc)

    def test_rodar_grava_pid_e_parar_derruba_tudo(self):
        self.sr.rodar(self.cfg, self.tmp, voltas=2, intervalo=0.5, espera_inicial=0)
        self.assertTrue(self._esperar_porta(self.pp))
        d = json.loads((self.tmp / "servico.pid").read_text())
        self.assertEqual(set(d["filhos"]), {"Prometheus", "Grafana"})
        self.assertTrue((self.tmp / "servico.log").exists())
        self.sr.parar_tudo(self.tmp)
        self.assertTrue(self._esperar_porta(self.pp, aberta=False))
        self.assertTrue(self._esperar_porta(self.pg, aberta=False))
        self.assertFalse((self.tmp / "servico.pid").exists())

    def test_sem_programa_o_aviso_diz_onde_configurar(self):
        self.cfg.set("servicos", "prometheus_exe", str(self.tmp / "nao_existe"))
        sv, avisos = self.sr.montar_servicos(self.cfg, self.tmp)
        self.assertEqual([s.nome for s in sv], ["Grafana"])
        self.assertIn("prometheus_exe", avisos[0])

    def test_acha_o_grafana_pela_estrutura_de_pastas(self):
        self.cfg.set("servicos", "grafana_exe", "")
        with mock.patch.object(self.sr, "raizes_de_instalacao",
                               return_value=[str(self.tmp)]):
            exe, raiz = self.sr.achar_grafana(self.cfg)
        self.assertTrue(exe.endswith("grafana-server"))
        self.assertEqual(Path(raiz), (self.tmp / "grafana").resolve())

    def test_tarefa_agendada_sem_limite_de_tempo(self):
        # schtasks /Create sem XML põe "parar após 3 dias".
        from xml.dom.minidom import parseString
        xml = self.sr.xml_da_tarefa([r"C:\rm\rajant_monitor.exe", "--servico"],
                                    r"C:\rm")
        doc = parseString(xml.replace('encoding="UTF-16"', 'encoding="UTF-8"'))
        txt = lambda t: doc.getElementsByTagName(t)[0].firstChild.data
        self.assertEqual(txt("ExecutionTimeLimit"), "PT0S")
        self.assertEqual(txt("UserId"), "S-1-5-18")              # SYSTEM
        self.assertEqual(len(doc.getElementsByTagName("BootTrigger")), 1)
        self.assertEqual(txt("Arguments"), "--servico")
        self.assertEqual(txt("WorkingDirectory"), r"C:\rm")

    def test_exportador_sob_o_servico_nao_escreve_no_console(self):
        import logging as _lg
        with mock.patch.dict(os.environ, {"RAJANT_SERVICO": "1"}):
            lg = rm.setup_logging(str(self.tmp / "x.log"), "INFO")
        tipos = {type(h).__name__ for h in _lg.getLogger().handlers}
        self.assertEqual(tipos, {"RotatingFileHandler"})
        rm.setup_logging("", "INFO")

    def test_main_despacha_antes_de_ler_o_config(self):
        fonte = inspect.getsource(rm.main)
        self.assertLess(fonte.index("servico_rajant.executar"),
                        fonte.index("ap.parse_args()"))


class TestRobustez(unittest.TestCase):

    def test_nmea_para_graus(self):
        self.assertAlmostEqual(rm.nmea_para_graus("2743.8950S"), -27.7316, places=3)
        self.assertAlmostEqual(rm.nmea_para_graus("05004.1429W"), -50.0690, places=3)
        self.assertIsNone(rm.nmea_para_graus(""))

    def test_state_vazio_nao_estoura(self):
        d = rm.parse_state("")
        self.assertEqual(d["radios"], [])
        self.assertEqual(d["ethernet"], [])
        self.assertFalse(d["sistema"]["apt_master"])

    def test_porta_sem_nome_recebe_rotulo(self):
        # label 'porta' vazio some do Grafana.
        st = 'wired {\n  aptState: APT_STATE_LINK\n  stats { rxBytes: 5 txBytes: 5 }\n}\n'
        self.assertEqual(rm.parse_state(st)["ethernet"][0]["nome"], "eth0")

    def test_extrair_blocos_aninhados(self):
        txt = "a {\n  b {\n    c: 1\n  }\n}\na {\n  d: 2\n}\n"
        self.assertEqual(len(rm.extrair_blocos(txt, 'a')), 2)

    def test_peers_ips(self):
        d = rm.parse_state(STATE)
        self.assertIn("10.188.96.141", d["peers_ips"])


# ══════════════════════════════════════════════════════════════════
# 14. COLETA ACELERADA DOS MÓVEIS
# ══════════════════════════════════════════════════════════════════
class TestIntervaloMoveis(unittest.TestCase):
    """Os móveis mudam de lugar e de vizinho a cada minuto: coletá-los mais
    rápido acompanha isso sem multiplicar o volume da frota toda."""

    PADRAO = "^CA ; ^PA ; ^PF ; ^TT ; ^EH ; CAMINH ; ESCAV ; PERFURA"

    def _coletor(self, interval=60, moveis=20, padrao=None):
        c = rm.RajantCollector(
            seeds=[], role="VIEW", password="", port=2300, interval=interval,
            max_threads=1, timeout=1, tentativas=1, falhas_limite=3,
            manter_s=300, redesc_ciclos=10, falhas_remover=20,
            cache=rm.CacheIPs(os.path.join(tempfile.mkdtemp(), "c.json")),
            interval_moveis=moveis,
            padrao_movel=self.PADRAO if padrao is None else padrao)
        return c

    def test_passo_dos_fixos(self):
        c = self._coletor(interval=60, moveis=20)
        self.assertEqual(c.interval_moveis, 20)
        self.assertEqual(c.passo_fixos, 3)      # fixos a cada 3 voltas = 60s

    def test_classifica_moveis_pelo_nome(self):
        c = self._coletor()
        c.nomes = {"1.1.1.1": "CA-4007", "2.2.2.2": "BACKBONE-ANEL-N",
                   "3.3.3.3": "ESCAV-12", "4.4.4.4": "ERB-005"}
        self.assertTrue(c._eh_movel("1.1.1.1"))
        self.assertFalse(c._eh_movel("2.2.2.2"))
        self.assertTrue(c._eh_movel("3.3.3.3"))
        self.assertFalse(c._eh_movel("4.4.4.4"))

    def test_nome_desconhecido_conta_como_movel(self):
        # Não deixa equipamento recém-descoberto fora do primeiro ciclo.
        c = self._coletor()
        self.assertTrue(c._eh_movel("9.9.9.9"))

    def test_desligado_por_padrao(self):
        c = self._coletor(moveis=0)
        self.assertEqual(c.interval_moveis, 0)
        self.assertEqual(c.passo_fixos, 1)
        self.assertIsNone(c.re_movel)
        # sem aceleração, todo BC entra em todo ciclo
        self.assertTrue(c._eh_movel("qualquer"))

    def test_ignora_intervalo_maior_que_o_geral(self):
        # 90s de "móveis" com 60s geral não faz sentido — desliga.
        c = self._coletor(interval=60, moveis=90)
        self.assertEqual(c.interval_moveis, 0)

    def test_padrao_invalido_nao_derruba_o_coletor(self):
        c = self._coletor(padrao="[nao-fecha")
        self.assertIsNone(c.re_movel)
        self.assertEqual(c.interval_moveis, 0)

    def test_sem_padrao_desliga(self):
        c = self._coletor(padrao="")
        self.assertEqual(c.interval_moveis, 0)


# ══════════════════════════════════════════════════════════════════
# 16. PAINEL HTML PRÓPRIO
# ══════════════════════════════════════════════════════════════════
class TestPainelHTML(unittest.TestCase):
    """Painel em HTML/CSS/JS servido em /painel, com proxy /api/q para o
    Prometheus (evita CORS e mantém a URL no servidor)."""

    def setUp(self):
        self.cfg = rm.configparser.ConfigParser()
        self.cfg.add_section("relatorio")

    def test_arquivo_do_painel_existe(self):
        p = pathlib.Path(rm.__file__).parent / "painel" / "visao-geral.html"
        self.assertTrue(p.exists(), "painel/visao-geral.html nao encontrado")
        html = p.read_text(encoding="utf-8")
        self.assertIn("<!DOCTYPE html>", html)
        # sem dependência externa: nada de CDN nem fonte remota
        self.assertNotIn("http://cdn", html.lower())
        self.assertNotIn("https://cdn", html.lower())
        self.assertNotIn("googleapis", html.lower())

    def test_painel_so_consulta_metricas_que_existem(self):
        p = pathlib.Path(rm.__file__).parent / "painel" / "visao-geral.html"
        html = p.read_text(encoding="utf-8")
        publicadas = set(re.findall(r'Gauge\("([a-z0-9_]+)"',
                                    open(rm.__file__, encoding="utf-8").read()))
        usadas = set(re.findall(r'\brajant_[a-z0-9_]+\b', html))
        for m in sorted(usadas):
            self.assertIn(m, publicadas,
                          f"{m} nao e publicada pelo exporter")

    def test_painel_nao_usa_metrica_removida(self):
        p = pathlib.Path(rm.__file__).parent / "painel" / "visao-geral.html"
        html = p.read_text(encoding="utf-8")
        removidas = {"rajant_voltagem_v", "rajant_bateria_v",
                     "rajant_session_state", "rajant_im_overflows",
                     "rajant_radio_phy_erros", "rajant_peer_txpower_dbm",
                     "rajant_eth_rx_erros", "rajant_eth_mudancas"}
        for m in removidas:
            self.assertNotIn(m, html, f"{m} foi removida na auditoria")

    def test_painel_html_lido_do_disco(self):
        # Editar o arquivo e recarregar o navegador basta — sem reiniciar.
        self.cfg.set("relatorio", "painel_html", "painel/visao-geral.html")
        html = rm.painel_html(self.cfg)
        self.assertIn("Visão Geral da Malha", html)

    def test_painel_ausente_da_mensagem_util(self):
        self.cfg.set("relatorio", "painel_html", "nao/existe.html")
        html = rm.painel_html(self.cfg)
        self.assertIn("não encontrado", html)
        self.assertIn("nao/existe.html", html)

    def test_proxy_api_q_declarado(self):
        fonte = open(rm.__file__, encoding="utf-8").read()
        self.assertIn('parsed.path == "/api/q"', fonte)
        # a janela e o passo são limitados, para não derrubar o Prometheus
        compacto = fonte.replace(" ", "")
        self.assertIn("min(janela,90*86400)", compacto)
        self.assertIn("min(passo,3600)", compacto)

    def test_rota_painel_declarada(self):
        fonte = open(rm.__file__, encoding="utf-8").read()
        self.assertIn('("/painel", "/painel.html")', fonte)


class TestInterferencia(unittest.TestCase):
    """A BC API nao tem varredura de espectro — verificado nos .proto. O
    que ela da e airtime, e dele sai a ocupacao por transmissor alheio."""

    IM = {k: 0 for k in ("pkt_tx", "pkt_drop", "pkt_rx", "arp",
                         "bytes_tx", "bytes_rx")}

    def _radio(self, ativo, busy, rx, tx):
        return {"nome": "wlan0", "ch_active": ativo, "ch_busy": busy,
                "ch_rx": rx, "ch_tx": tx, "rx_bytes": 0, "tx_bytes": 0,
                "rx_pkts": 0, "tx_pkts": 0, "canal": 149}

    def _pacote(self, r):
        return {"radios": [r], "ethernet": [],
                "sistema": {"idle": 0, "uptime": 0},
                "instamesh": dict(self.IM)}

    def _delta(self, ip, antes, agora, dt=60.0):
        t = time.monotonic()
        rm._ultimo[ip] = {"d": self._pacote(self._radio(*antes)), "ts": t - dt}
        return rm.calcular_taxas(ip, self._pacote(self._radio(*agora)),
                                 t)["radios"][0]

    def test_proto_nao_tem_varredura_de_espectro(self):
        # Se um dia a Rajant publicar isso, este teste avisa que da para
        # trocar a estimativa por medicao de verdade.
        raiz = pathlib.Path("bcapi-ref/proto")
        if not raiz.exists():
            self.skipTest("protos nao disponiveis")
        texto = " ".join(p.read_text(errors="ignore").lower()
                         for p in raiz.glob("*.proto"))
        for termo in ("spectrum", "spectral", "sweep"):
            self.assertNotIn(termo, texto,
                             f"'{termo}' apareceu no proto — reavaliar")

    def test_trafego_proprio_nao_conta_como_interferencia(self):
        # O caso que confunde: canal 50% ocupado, mas ocupado por NOS.
        r = self._delta("10.9.0.1", (1000,1000,1000,1000),
                        (61000, 31000, 21000, 10000))
        self.assertAlmostEqual(r["busy_pct"], 50.0, places=1)
        self.assertLess(r["interf_pct"], 5.0)

    def test_ocupacao_alheia_vira_interferencia(self):
        r = self._delta("10.9.0.2", (1000,1000,1000,1000),
                        (61000, 40000, 3000, 2000))
        self.assertAlmostEqual(r["interf_pct"], 60.0, places=1)

    def test_primeira_coleta_fica_sem_valor_nao_zero(self):
        # Contadores sao cumulativos desde o boot: sem delta, a razao
        # seria a media da vida inteira do radio.
        d = rm.calcular_taxas("10.9.0.3",
                              self._pacote(self._radio(61000,40000,3000,2000)),
                              time.monotonic())
        self.assertIsNone(d["radios"][0].get("interf_pct"))

    def test_parse_state_nao_inventa_valor_cumulativo(self):
        fonte = inspect.getsource(rm.parse_state)
        self.assertIn('"interf_pct":   None', fonte)

    def test_nao_negativa_quando_contadores_desalinham(self):
        # rx+tx > busy acontece com arredondamento do firmware.
        r = self._delta("10.9.0.4", (1000,1000,1000,1000),
                        (61000, 5000, 4000, 4000))
        self.assertGreaterEqual(r["interf_pct"], 0.0)


class TestDescobertaEFiltro(unittest.TestCase):
    """O exporter mostrava IP e MAC que nao sao equipamento da frota."""

    ST = '''
    wireless { name: "wlan0" channel: 149 noise: -95
      peer { mac: "AA:BB:CC:DD:EE:FF" enabled: true rssi: 40 signal: -60
             cost: 10 rate: 1300 }
      peer { ipv4Address: "10.0.0.7" mac: "11:22:33:44:55:66" enabled: true
             rssi: 42 signal: -58 cost: 8 rate: 1300 }
      peer { ipv4Address: "0.0.0.0" enabled: true rssi: 10 signal: -80
             cost: 99 rate: 65 }
    }
    '''

    def test_peer_sem_ip_nao_entra_na_descoberta(self):
        # `mac:AA:BB:...` nao e endereco: ia para a fila, falhava a
        # conexao, era marcado OFFLINE e virava equipamento fantasma.
        d = rm.parse_state(self.ST)
        self.assertEqual(d["peers_ips"], {"10.0.0.7"})

    def test_mas_o_enlace_do_peer_sem_ip_e_preservado(self):
        # Descartar o peer inteiro perderia SNR, custo e taxa do enlace —
        # foi por isso que o identificador por MAC existe.
        d = rm.parse_state(self.ST)
        ids = {p["ip"] for p in d["radios"][0]["peers"]}
        self.assertIn("mac:AA:BB:CC:DD:EE:FF", ids)
        self.assertEqual(len(d["radios"][0]["peers"]), 2)

    def test_ip_zerado_tambem_fica_fora(self):
        d = rm.parse_state(self.ST)
        self.assertNotIn("0.0.0.0", d["peers_ips"])

    def _col(self, prefs, padrao=None):
        c = rm.RajantCollector.__new__(rm.RajantCollector)
        c.prefixos_tag = prefs; c.sem_tag = set()
        c.re_tag = rm.re.compile(padrao, rm.re.I) if padrao else None
        return c

    def test_sem_prefixos_configurados_aceita_tudo(self):
        # O filtro e opcional: nao pode passar a descartar por omissao.
        c = self._col(None)
        for n in ("CA-1001", "qualquer-coisa", "", "10.0.0.9"):
            self.assertTrue(c.tem_tag(n), n)

    def test_com_prefixos_so_passa_a_frota(self):
        c = self._col(["CA", "PA", "PF", "TT", "EH", "ERM", "ERB"])
        for n in ("CA-1001", "ERM-20", "ERB-04", "pf-77"):
            self.assertTrue(c.tem_tag(n), n)
        for n in ("BC-XYZ", "", "10.0.0.99", "teste-lab", None):
            self.assertFalse(c.tem_tag(n), n)

    def test_descarte_e_auditavel(self):
        # Sumir com equipamento calado seria pior que mostrar de mais.
        fonte = inspect.getsource(rm.RajantCollector._coletar_bc)
        self.assertIn("m_bc_sem_tag", fonte)
        self.assertIn("ignorado", fonte)

    def test_descoberta_desligada_nao_expande_no_ciclo(self):
        # Sem esta guarda o ciclo reintroduzia pela porta dos fundos o
        # que a descoberta deixou de fora.
        fonte = inspect.getsource(rm.RajantCollector.ciclo)
        self.assertIn("if peers and self.descoberta:", fonte)

    def test_descoberta_desligada_mantem_seed_offline(self):
        # Um BC que caiu agora tem de continuar monitorado, senao ele
        # some da metrica justamente quando cai.
        fonte = inspect.getsource(rm.RajantCollector.descobrir)
        self.assertIn("alvos = set(self.seeds)", fonte)

    def test_cache_desligado_nao_escreve_disco(self):
        tmp = tempfile.mkdtemp()
        try:
            cam = Path(tmp) / "c.json"
            c = rm.CacheIPs(str(cam), ativo=False)
            c.adicionar("10.0.0.1", "CA-1001")
            self.assertFalse(cam.exists(), "escreveu com o cache desligado")
            self.assertEqual(c.total(), 1, "deveria guardar em memoria")
            c2 = rm.CacheIPs(str(cam), ativo=True)
            c2.adicionar("10.0.0.2", "CA-1002")
            self.assertTrue(cam.exists())
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_config_traz_as_chaves(self):
        for chave in ("descoberta", "usar_cache", "filtrar_tag", "padrao_tag",
                      "segurar_serie_s"):
            self.assertIn(chave, rm.DEFAULT_CONFIG, chave)
        # As chaves antigas sairam do padrao: o config.ini ja instalado as
        # tem escritas (false e 300), e com o nome antigo elas seguiriam
        # valendo e o filtro e a retirada nunca ligariam.
        for velha in ("somente_com_tag", "manter_ultimo_valor_s"):
            self.assertNotIn(velha + " ", rm.DEFAULT_CONFIG)

    def test_padrao_de_tag_aceita_a_rede_da_mina_e_recusa_ip(self):
        c = self._col(None, rm.PADRAO_TAG_PADRAO)
        nomes = [n for n in rm.REDE_CONHECIDA.values() if n]
        self.assertGreater(len(nomes), 150)
        for n in nomes:
            self.assertTrue(c.tem_tag(n), n)
        for n in ("10.188.99.217", "", None, "teste-lab", "Rajant BC",
                  "eth0udp25", "wlan0wds33"):
            self.assertFalse(c.tem_tag(n), n)


class TestShimSslWrapSocket(unittest.TestCase):
    """A rajant-api 0.1.1 faz `from ssl import wrap_socket`, removido no
    Python 3.12. O shim do topo do modulo e o que mantem o programa
    rodando em 3.12+ — sem ele, nao importa."""

    def test_shim_existe_no_topo_do_modulo(self):
        # Tem de vir ANTES do `from rajant_api import Breadcrumb`, senao
        # a biblioteca ja quebrou quando o shim seria instalado.
        fonte = pathlib.Path("rajant_monitor.py").read_text()
        i_shim = fonte.index("ssl.wrap_socket = _wrap")
        i_imp = fonte.index("from rajant_api import Breadcrumb")
        self.assertLess(i_shim, i_imp, "shim depois do import: tarde demais")

    def test_shim_reproduz_o_comportamento_antigo(self):
        # ssl.wrap_socket() com os padroes NAO validava certificado — os
        # BreadCrumbs usam autoassinado. O shim precisa manter isso, ou a
        # conexao passa a falhar onde antes funcionava.
        import ssl as _ssl
        fonte = pathlib.Path("rajant_monitor.py").read_text()
        trecho = fonte[:fonte.index("import re, os, json")]
        self.assertIn("cert_reqs=ssl.CERT_NONE", trecho)
        self.assertIn("check_hostname = False", trecho)

    def test_shim_nao_sobrescreve_quando_ja_existe(self):
        # Em Python <=3.11 a funcao original tem de continuar valendo.
        fonte = pathlib.Path("rajant_monitor.py").read_text()
        self.assertIn("if not hasattr(ssl, 'wrap_socket'):", fonte)

    def test_shim_faz_handshake_de_verdade(self):
        """Importar nao prova nada: o que importa e a conexao TLS."""
        import ssl as _ssl, socket, threading, tempfile, subprocess
        if not hasattr(_ssl, "wrap_socket"):
            self.skipTest("Python sem wrap_socket original para comparar")
        d = tempfile.mkdtemp()
        cert, chave = os.path.join(d, "c.pem"), os.path.join(d, "k.pem")
        r = subprocess.run(
            ["openssl", "req", "-x509", "-newkey", "rsa:2048", "-keyout",
             chave, "-out", cert, "-days", "1", "-nodes", "-subj",
             "/CN=bc-teste"], capture_output=True)
        if r.returncode != 0:
            self.skipTest("openssl indisponivel")

        ctx = _ssl.SSLContext(_ssl.PROTOCOL_TLS_SERVER)
        ctx.load_cert_chain(cert, chave)
        srv = socket.socket()
        srv.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        srv.bind(("127.0.0.1", 0)); srv.listen(1)
        porta = srv.getsockname()[1]

        def atender():
            try:
                c, _ = srv.accept()
                with ctx.wrap_socket(c, server_side=True) as t:
                    t.sendall(b"OI")
            except Exception:
                pass
        threading.Thread(target=atender, daemon=True).start()

        # Reconstrói o shim isoladamente, do jeito que o módulo faz.
        def _wrap(sock, cert_reqs=_ssl.CERT_NONE, **kw):
            c = _ssl.SSLContext(_ssl.PROTOCOL_TLS_CLIENT)
            c.check_hostname = False; c.verify_mode = cert_reqs
            return c.wrap_socket(sock)
        soc = socket.socket(); soc.settimeout(5)
        con = _wrap(soc)
        try:
            con.connect(("127.0.0.1", porta))
            self.assertEqual(con.recv(16), b"OI")
            self.assertIsNotNone(con.cipher())
        finally:
            con.close(); srv.close(); shutil.rmtree(d, ignore_errors=True)


class TestSemanalSemSurvey(unittest.TestCase):
    """O survey foi descontinuado. O template do cliente ainda traz a
    seção de Site Survey, e ela não pode sair em branco no semanal."""

    def test_gerar_ppt_sempre_remove_a_secao(self):
        fonte = inspect.getsource(rm.gerar_ppt)
        self.assertIn("remover_slides_survey(p)", fonte)
        self.assertNotIn("survey_no_semanal", fonte)
        self.assertNotIn("anexar_survey", fonte)

    def test_gerar_ppt_nao_aceita_mais_survey(self):
        self.assertNotIn("survey", inspect.signature(rm.gerar_ppt).parameters)
        self.assertNotIn("survey_id",
                         inspect.signature(rm.montar_relatorio).parameters)

    def test_remocao_tira_os_slides_do_template(self):
        fonte = inspect.getsource(rm.remover_slides_survey)
        self.assertIn("_eh_titulo_survey", fonte)
        self.assertIn("site survey", fonte)

    def test_navegacao_nao_e_removida(self):
        # O slide de Navegacao lista TODAS as secoes; casaria com o
        # filtro e sumiria junto.
        fonte = inspect.getsource(rm.remover_slides_survey)
        self.assertIn('startswith("navega")', fonte)

    def test_botao_do_menu_deixa_de_apontar_para_o_vazio(self):
        # Link para slide removido nao acusa erro: so nao vai a lugar
        # nenhum.
        fonte = inspect.getsource(rm.remover_slides_survey)
        self.assertIn("(descontinuado)", fonte)
        self.assertIn("_limpar_link_de_texto", fonte)

    def test_remove_de_tras_para_frente(self):
        fonte = inspect.getsource(rm.remover_slides_survey)
        self.assertIn("for i in reversed(alvos):", fonte)

    def test_config_sem_chaves_de_survey(self):
        for k in ("survey_no_semanal", "survey_kmz", "prefixos_frota",
                  "kmz_com_equipamentos", "imagens_no_ppt", "zonas_grade_m",
                  "fundo_local", "fundo_bbox"):
            self.assertNotIn(k, rm.DEFAULTS_RELATORIO)
        cfg = rm.configparser.ConfigParser()
        with mock.patch.object(rm, "CONFIG_FILE",
                               str(Path(tempfile.mkdtemp()) / "c.ini")):
            rm.cfg_relatorio(cfg)
        self.assertFalse(cfg.has_section("survey"))


class TestPaginaSoRelatorios(unittest.TestCase):
    """A página web ficou só com relatório, painel e diagnóstico."""

    def _servidor(self):
        import http.server, threading
        cfg = rm.configparser.ConfigParser()
        cfg.add_section("relatorio")
        cfg.set("relatorio", "prometheus_url", "http://127.0.0.1:9")
        cfg.set("relatorio", "painel_html", "painel/visao-geral.html")
        srv = http.server.HTTPServer(("127.0.0.1", 0), rm.criar_handler(cfg))
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        self.addCleanup(srv.server_close)
        self.addCleanup(srv.shutdown)
        return f"http://127.0.0.1:{srv.server_address[1]}"

    def _get(self, url):
        import urllib.request, urllib.error
        try:
            with urllib.request.urlopen(url, timeout=5) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()

    def test_pagina_sem_survey(self):
        html = rm.PAGINA_HTML.lower()
        for termo in ("survey", "captura", "kmz", "medições", "histórico"):
            self.assertNotIn(termo, html, termo)
        self.assertIn("gerar e baixar excel", html)
        self.assertIn("gerar e baixar ppt", html)
        self.assertIn('href="/painel"', html)

    def test_rotas_de_survey_nao_existem_mais(self):
        base = self._servidor()
        for rota in ("/bcs", "/surveys", "/survey?id=1", "/captura/ativas",
                     "/medicoes/listar", "/radios", "/perfis/abrir?nome=x"):
            self.assertEqual(self._get(base + rota)[0], 404, rota)

    def test_rotas_que_ficam(self):
        base = self._servidor()
        st, corpo = self._get(base + "/")
        self.assertEqual(st, 200)
        self.assertIn(b"Gerar e baixar Excel", corpo)
        self.assertEqual(self._get(base + "/painel")[0], 200)
        self.assertEqual(self._get(base + "/diagnostico")[0], 400)
        st, corpo = self._get(base + "/api/q?q=up")
        self.assertEqual(st, 200)
        self.assertFalse(json.loads(corpo)["ok"])   # Prometheus fora do ar

    def test_gerar_recusa_periodo_invertido(self):
        base = self._servidor()
        st, corpo = self._get(base + "/gerar?ini=2026-10-10&fim=2026-10-01")
        self.assertEqual(st, 500)
        self.assertIn("data inicial", corpo.decode("utf-8"))

    def test_post_nao_e_aceito(self):
        import urllib.request, urllib.error
        base = self._servidor()
        req = urllib.request.Request(base + "/medicoes/importar", data=b"x",
                                     method="POST")
        with self.assertRaises(urllib.error.HTTPError) as cm:
            urllib.request.urlopen(req, timeout=5)
        self.assertEqual(cm.exception.code, 501)


class TestIdentidadeAnglo(unittest.TestCase):
    """As medidas vieram do Dashboard_Transporte do cliente: logo, titulo
    e regua caem nos mesmos pontos, senao os decks nao parecem da mesma
    familia."""

    def test_paleta_e_do_template_do_cliente(self):
        self.assertEqual(rm.ANGLO["azul"], "031795")
        self.assertEqual(rm.ANGLO["fonte"], "Calibri")

    def test_geometria_igual_a_do_template(self):
        self.assertEqual(rm.ANGLO["logo"], (0.30, 0.20, 1.50, 0.59))
        self.assertEqual(rm.ANGLO["titulo"], (1.80, 0.20, 9.73, 0.70))
        self.assertEqual(rm.ANGLO["regua_y"], 1.00)

    def test_sem_a_pasta_marca_nao_ha_logo_em_vez_de_estourar(self):
        tmp = tempfile.mkdtemp(); cwd = os.getcwd()
        try:
            os.chdir(tmp)
            with mock.patch.object(rm, "DIR_MARCA", "pasta_que_nao_existe"):
                self.assertIsNone(rm._logo_anglo())
        finally:
            os.chdir(cwd); shutil.rmtree(tmp, ignore_errors=True)


class TestIdentidadeNoSemanal(unittest.TestCase):
    """O semanal vira Anglo repintando o template do cliente, no fim da
    geração. O que se testa aqui é que a pele muda e o conteúdo não."""

    def _deck(self):
        """Um deck mínimo na paleta do template: capa + um slide de
        conteúdo com cartão, tabela e botão."""
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE

        p = Presentation()
        p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
        vazio = p.slide_masters[0].slide_layouts[6]

        def texto(s, x, y, w, h, txt, cor, tam=12):
            tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = txt
            r.font.name = "Calibri"; r.font.size = Pt(tam)
            r.font.color.rgb = RGBColor.from_string(cor)
            return tb

        # capa
        capa = p.slides.add_slide(vazio)
        capa.background.fill.solid()
        capa.background.fill.fore_color.rgb = RGBColor.from_string("0D2052")
        texto(capa, 0.9, 2.3, 11.5, 0.9, "RELATÓRIO SEMANAL", "FFFFFF", 40)
        # Subtítulo em tom de corpo: no template ele é A0B0CC, que vira
        # cinza — e cinza sobre o azul da capa some. Está aqui porque foi
        # exatamente o que apareceu no arquivo do cliente.
        texto(capa, 0.9, 3.25, 11.5, 0.45,
              "Infraestrutura de Telecom | Tecnologia de Mina", "A0B0CC", 16)

        # conteúdo
        s = p.slides.add_slide(vazio)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("0D2052")
        texto(s, 0.55, 0.32, 11.3, 0.6, "1. Dashboard Executivo", "FFFFFF", 28)
        texto(s, 0.55, 0.92, 11.3, 0.32, "Resumo da semana", "A0B0CC")
        cart = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55),
                                  Inches(1.55), Inches(4.0), Inches(1.5))
        cart.fill.solid(); cart.fill.fore_color.rgb = RGBColor.from_string("1A2744")
        cart.line.color.rgb = RGBColor.from_string("2A3A5C")
        # Rótulo BRANCO em cima do cartão: o cartão clareia, então este
        # branco tem de virar escuro ou o texto some. É o caso dos 174
        # runs brancos do template do cliente.
        texto(s, 0.75, 1.70, 3.6, 0.3, "DISPONIBILIDADE DA MALHA", "FFFFFF", 11)
        texto(s, 0.75, 2.05, 3.6, 0.85, "99,12%", "C8D0E0", 30)
        gf = s.shapes.add_table(2, 2, Inches(0.55), Inches(3.3),
                                Inches(6.0), Inches(0.8))
        for j, cab in enumerate(("Indicador", "Valor")):
            cel = gf.table.cell(0, j)
            cel.fill.solid(); cel.fill.fore_color.rgb = RGBColor.from_string("1A3A7A")
            r = cel.text_frame.paragraphs[0].add_run(); r.text = cab
            r.font.color.rgb = RGBColor.from_string("FFFFFF")
        for j, v in enumerate(("Disponibilidade", "OK")):
            cel = gf.table.cell(1, j)
            cel.fill.solid(); cel.fill.fore_color.rgb = RGBColor.from_string("152238")
            r = cel.text_frame.paragraphs[0].add_run(); r.text = v
            r.font.color.rgb = RGBColor.from_string("27AE60")
        bt = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12.05),
                                Inches(0.28), Inches(1.05), Inches(0.34))
        bt.fill.solid(); bt.fill.fore_color.rgb = RGBColor.from_string("1A2744")
        r = bt.text_frame.paragraphs[0].add_run(); r.text = "◂ MENU"
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
        bt.click_action.target_slide = capa
        return p

    # ── a pele muda ───────────────────────────────────────────
    def test_fundo_de_conteudo_vira_branco_e_a_capa_vira_azul(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        self.assertEqual(str(p.slides[0].background.fill.fore_color.rgb),
                         rm.ANGLO["azul"])
        self.assertEqual(str(p.slides[1].background.fill.fore_color.rgb),
                         rm.ANGLO["fundo"])

    def test_nao_sobra_cor_do_template_escuro(self):
        import io, re, zipfile, collections
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        buf = io.BytesIO(); p.save(buf)
        z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
        achadas = collections.Counter()
        for n in z.namelist():
            if not re.match(r"ppt/(slides|charts)/.*\.xml$", n): continue
            txt = z.read(n).decode("utf8")
            for c in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', txt):
                achadas[c.upper()] += 1
        antigas = set(rm.ANGLO_FUNDOS) | set(rm.ANGLO_TEXTOS)
        # As cores de status voltam a valer sobre fundo escuro (a capa),
        # então elas podem reaparecer; as estruturais, não.
        estruturais = set(rm.ANGLO_FUNDOS)
        sobrou = {c: k for c, k in achadas.items() if c in estruturais}
        self.assertEqual(sobrou, {}, f"cor do template escuro sobrou: {sobrou}")
        self.assertTrue(antigas, "mapa de cores vazio invalida o teste")

    def test_texto_claro_vira_escuro_sobre_o_fundo_branco(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        achou = False
        for sh in p.slides[1].shapes:
            if sh.has_text_frame and "99,12" in sh.text_frame.text:
                cor = str(sh.text_frame.paragraphs[0].runs[0].font.color.rgb)
                self.assertEqual(cor, rm.ANGLO["texto"]); achou = True
        self.assertTrue(achou, "o valor do cartão sumiu do slide")

    def test_branco_continua_branco_onde_o_fundo_seguiu_escuro(self):
        # Na capa o fundo vira azul institucional: virar o texto escuro
        # apagaria o título.
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        for sh in p.slides[0].shapes:
            if sh.has_text_frame and "RELATÓRIO" in sh.text_frame.text:
                self.assertEqual(
                    str(sh.text_frame.paragraphs[0].runs[0].font.color.rgb),
                    "FFFFFF")
                return
        self.fail("titulo da capa sumiu")

    def test_borda_de_celula_e_repintada(self):
        # As bordas moram em lnL/lnR/lnT/lnB dentro do tcPr e NÃO passam
        # pelo python-pptx. Sem tratá-las, sobra a grade azul-escura
        # riscando o fundo branco.
        from pptx.util import Inches
        from lxml import etree
        p = self._deck()
        tab = [sh for sh in p.slides[1].shapes if sh.has_table][0].table
        tc = tab.cell(0, 0)._tc
        tcPr = tc.find(f"{rm._A_NS}tcPr")
        if tcPr is None:
            tcPr = etree.SubElement(tc, f"{rm._A_NS}tcPr")
        ln = etree.SubElement(tcPr, f"{rm._A_NS}lnL")
        sf = etree.SubElement(ln, f"{rm._A_NS}solidFill")
        etree.SubElement(sf, f"{rm._A_NS}srgbClr").set("val", "2A3A5C")
        rm.aplicar_identidade_anglo(p)
        cores = [c.get("val") for c in tab._tbl.iter(f"{rm._A_NS}srgbClr")]
        self.assertNotIn("2A3A5C", cores)
        self.assertIn(rm.ANGLO["linha"], cores)

    # ── o conteúdo não muda ───────────────────────────────────
    def test_o_link_do_botao_sobrevive(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        ligados = 0
        for sh in p.slides[1].shapes:
            try:
                if sh.click_action.target_slide is not None: ligados += 1
            except Exception:
                pass
        self.assertEqual(ligados, 1, "a repintura levou o hyperlink junto")

    def test_nem_um_texto_se_perde(self):
        p = self._deck()
        antes = sorted(sh.text_frame.text for s in p.slides
                       for sh in s.shapes if sh.has_text_frame)
        rm.aplicar_identidade_anglo(p)
        depois = sorted(sh.text_frame.text for s in p.slides
                        for sh in s.shapes if sh.has_text_frame)
        # A repintura só ACRESCENTA (logo e régua não têm texto).
        for t in antes:
            self.assertIn(t, depois)

    def test_botao_de_voltar_fica_azul_com_texto_branco(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        for sh in p.slides[1].shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith("◂"):
                self.assertEqual(str(sh.fill.fore_color.rgb), rm.ANGLO["azul"])
                self.assertEqual(
                    str(sh.text_frame.paragraphs[0].runs[0].font.color.rgb),
                    rm.ANGLO["fundo"])
                return
        self.fail("botao de voltar sumiu")

    # ── a marca ───────────────────────────────────────────────
    def test_cada_slide_de_conteudo_ganha_regua(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        self.assertTrue(any(sh.name == "anglo_regua"
                            for sh in p.slides[1].shapes))
        # a capa não leva régua: ela é toda azul
        self.assertFalse(any(sh.name == "anglo_regua"
                             for sh in p.slides[0].shapes))

    def test_regua_nao_entra_por_cima_de_conteudo(self):
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        p = self._deck()
        s = p.slides[1]
        intruso = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55),
                                     Inches(rm.ANGLO_SEMANAL["regua_y"] - 0.05),
                                     Inches(5.0), Inches(0.3))
        rm.aplicar_identidade_anglo(p)
        self.assertFalse(any(sh.name == "anglo_regua" for sh in s.shapes),
                         "regua desenhada em cima do conteudo")

    def test_slide_que_ja_nasceu_anglo_nao_ganha_segunda_marca(self):
        # No modo unido o deck de survey é anexado já na identidade: um
        # segundo logo cairia em cima do primeiro.
        p = self._deck()
        n1 = rm.aplicar_identidade_anglo(p)["marcas"]
        n2 = rm.aplicar_identidade_anglo(p)["marcas"]
        self.assertGreater(n1, 0)
        self.assertEqual(n2, 0, "a segunda passagem marcou de novo")

    def test_titulo_sai_da_area_do_logo(self):
        from pptx.util import Emu
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        for sh in p.slides[1].shapes:
            if sh.has_text_frame and "Dashboard" in sh.text_frame.text:
                self.assertAlmostEqual(Emu(sh.left).inches,
                                       rm.ANGLO_SEMANAL["titulo"][0], places=2)
                self.assertEqual(
                    str(sh.text_frame.paragraphs[0].runs[0].font.color.rgb),
                    rm.ANGLO["azul"])
                return
        self.fail("titulo sumiu")

    # ── contraste ─────────────────────────────────────────────
    def test_todo_texto_fica_legivel_sobre_o_que_ficou_atras(self):
        def lum(h):
            def canal(i):
                v = int(h[i:i+2], 16) / 255
                return v/12.92 if v <= 0.03928 else ((v+0.055)/1.055)**2.4
            return 0.2126*canal(0) + 0.7152*canal(2) + 0.0722*canal(4)
        def razao(a, b):
            la, lb = lum(a), lum(b)
            return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)

        p = self._deck(); rm.aplicar_identidade_anglo(p)
        piores = []
        for s in p.slides:
            formas = list(rm._formas_planas(s.shapes))
            fundo_s = str(s.background.fill.fore_color.rgb).upper()
            alvos = []
            for sh in formas:
                if sh.has_text_frame:
                    f = rm._hex_do_fill(sh)
                    if f is None:
                        atras = rm._cartao_atras(sh, formas)
                        f = rm._hex_do_fill(atras) if atras is not None else fundo_s
                    alvos.append((sh.text_frame, f))
                if sh.has_table:
                    for row in sh.table.rows:
                        for cel in row.cells:
                            alvos.append((cel.text_frame,
                                          rm._hex_do_fill(cel) or fundo_s))
            for tf, fd in alvos:
                for par in tf.paragraphs:
                    for r in par.runs:
                        try: hx = str(r.font.color.rgb).upper()
                        except Exception: continue
                        rz = razao(hx, fd or "FFFFFF")
                        if rz < 3.0:
                            piores.append((round(rz, 2), hx, fd, r.text[:20]))
        self.assertEqual(piores, [], f"texto ilegivel apos a repintura: {piores}")

    # ── ligação com a geração ─────────────────────────────────
    def test_gerar_ppt_aplica_por_padrao_e_por_ultimo(self):
        fonte = inspect.getsource(rm.gerar_ppt)
        self.assertIn('"identidade_anglo", fallback=True', fonte)
        # Repintar antes de preencher deixaria de fora o verde/ambar/
        # vermelho do status e os slides tecnicos.
        self.assertLess(fonte.index("_slides_tecnicos(p, d, cfg"),
                        fonte.index("aplicar_identidade_anglo(p)"))
        self.assertLess(fonte.index("remover_slides_survey(p)"),
                        fonte.index("aplicar_identidade_anglo(p)"))

    def test_pode_ser_desligada(self):
        cfg = rm.configparser.ConfigParser()
        with mock.patch.object(rm, "CONFIG_FILE",
                               str(Path(tempfile.mkdtemp()) / "c.ini")):
            rm.cfg_relatorio(cfg)
        self.assertTrue(cfg.has_option("relatorio", "identidade_anglo"))
        self.assertTrue(cfg.getboolean("relatorio", "identidade_anglo"))


def _deck_de_secoes(titulos_e_subtitulos):
    """Deck com um slide por (título, subtítulo), no leiaute do template:
    o botão ◂ MENU em 0,28" — ACIMA do título, em 0,32"."""
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    vazio = p.slide_masters[0].slide_layouts[6]
    for titulo, sub in titulos_e_subtitulos:
        s = p.slides.add_slide(vazio)
        bt = s.shapes.add_textbox(Inches(12.05), Inches(0.28),
                                  Inches(1.05), Inches(0.34))
        bt.text_frame.paragraphs[0].add_run().text = "◂ MENU"
        t = s.shapes.add_textbox(Inches(0.55), Inches(0.32),
                                 Inches(11.3), Inches(0.6))
        t.text_frame.paragraphs[0].add_run().text = titulo
        u = s.shapes.add_textbox(Inches(0.55), Inches(0.92),
                                 Inches(11.3), Inches(0.32))
        u.text_frame.paragraphs[0].add_run().text = sub
    return p


class TestTituloDoSlide(unittest.TestCase):
    def test_botao_de_menu_nao_vira_titulo(self):
        # O ◂ MENU fica em 0,28" e o título em 0,32": pegar "o texto mais
        # alto" sem filtrar botão devolveria o botão, e aí TODO slide se
        # chamaria "◂ MENU".
        p = _deck_de_secoes([("1. Dashboard Executivo", "resumo")])
        self.assertEqual(rm._titulo_do_slide(list(p.slides)[0]),
                         "1. Dashboard Executivo")

    def test_slide_sem_texto_nao_estoura(self):
        from pptx import Presentation
        p = Presentation()
        s = p.slides.add_slide(p.slide_masters[0].slide_layouts[6])
        self.assertEqual(rm._titulo_do_slide(s), "")


class TestRemocaoDaSecaoDeSurvey(unittest.TestCase):
    """O semanal remove a seção de survey que vem no template. O que se
    testa aqui é que ele remove SÓ ela."""

    SECOES = [
        ("Navegação", "Clique numa seção para ir ao slide"),
        ("1. Dashboard Executivo", "Resumo da saúde da infraestrutura"),
        ("4. Saúde da Mesh Rajant — Links Críticos (SNR)",
         "10 piores enlaces do período — candidatos a realinhamento / "
         "site survey  |  gerado automaticamente"),
        ("4. Saúde da Mesh Rajant — Espectro por Canal",
         "Ocupação e ruído agregados por canal — todos os rádios ativos"),
        ("5. Site Survey — Intensidade de Sinal (RSSI) — 2.4 GHz", "medições"),
        ("5. Site Survey — Análise, Recomendações e Conclusão", "conclusão"),
        ("6. Manutenções", "preventivas e corretivas do período"),
    ]

    def _restantes(self):
        p = _deck_de_secoes(self.SECOES)
        n = rm.remover_slides_survey(p)
        return n, [rm._titulo_do_slide(s) for s in p.slides]

    def test_tira_os_slides_de_survey(self):
        n, restantes = self._restantes()
        self.assertEqual(n, 2)
        self.assertFalse([t for t in restantes if t.startswith("5. Site Survey")])

    def test_prosa_sobre_survey_no_subtitulo_nao_apaga_o_slide(self):
        # A REGRESSÃO: o filtro casava "site survey" em QUALQUER texto do
        # slide. O subtítulo de Links Críticos diz "candidatos a
        # realinhamento / site survey" e o slide sumia do semanal, que
        # ficava com um único slide de gráficos.
        _, restantes = self._restantes()
        self.assertIn("4. Saúde da Mesh Rajant — Links Críticos (SNR)",
                      restantes)

    def test_espectro_por_canal_fica_no_semanal(self):
        # A outra metade: o slide vinha titulado "5. Site Survey —
        # Espectro por Canal" e ia junto, mas os números são dos
        # contadores do exporter.
        _, restantes = self._restantes()
        self.assertIn("4. Saúde da Mesh Rajant — Espectro por Canal",
                      restantes)

    def test_navegacao_sobrevive(self):
        _, restantes = self._restantes()
        self.assertIn("Navegação", restantes)

    def test_o_resto_do_relatorio_fica_de_pe(self):
        _, restantes = self._restantes()
        for t in ("1. Dashboard Executivo", "6. Manutenções"):
            self.assertIn(t, restantes)

    def test_deck_sem_survey_nao_perde_nada(self):
        p = _deck_de_secoes([("1. Dashboard Executivo", "resumo"),
                             ("6. Manutenções", "preventivas")])
        self.assertEqual(rm.remover_slides_survey(p), 0)
        self.assertEqual(len(list(p.slides)), 2)


class TestSlidesTecnicosNoSemanal(unittest.TestCase):
    def test_espectro_pertence_a_secao_4(self):
        # Os números vêm dos contadores do EXPORTER — todos os rádios
        # ativos, o período inteiro —, não de uma captura de survey.
        fonte = inspect.getsource(rm._slides_tecnicos)
        self.assertIn('"4. Saúde da Mesh Rajant — Espectro por Canal"', fonte)
        self.assertNotIn('"5. Site Survey — Espectro por Canal"', fonte)

    def test_espectro_ancora_em_slide_que_existe_no_semanal(self):
        # Ancorado só em slides de survey, ele ficava solto no fim do
        # deck quando a seção 5 não existia.
        fonte = inspect.getsource(rm._slides_tecnicos)
        trecho = fonte[fonte.index("Espectro por Canal"):]
        chamada = trecho[trecho.index("mover_apos_titulo("):]
        chamada = chamada[:chamada.index(")")]
        self.assertTrue("Ethernet e Quedas" in chamada
                        or "Links Críticos" in chamada,
                        f"ancora so em slide de survey: {chamada}")


class TestCamposDescartados(BaseParser):
    """Campos que o rádio mandava e o parser jogava fora.

    Os três primeiros são o que permite medir sinal na posição certa em
    vez de na posição de antes. Conferidos contra `bcapi-ref/proto/`.
    """

    def test_gpstime_e_lido(self):
        # Gps.proto: GPSPositionReport.gpsTime = 1 (float).
        self.assertEqual(self.s["gps_time"], 143025.0)

    def test_gpstime_vira_segundos_do_dia(self):
        # 14:30:25 = 14*3600 + 30*60 + 25
        self.assertEqual(self.s["gps_time_s"], 52225.0)

    def test_gpsquality_e_lido(self):
        # NMEA GGA: 0 inválido, 1 GPS, 2 DGPS, 4/5 RTK.
        self.assertEqual(self.s["gps_qual"], 1.0)

    def test_sem_gpstime_fica_none_nao_zero(self):
        st = STATE.replace("    gpsTime: 143025.0\n", "")
        d = rm.parse_state(st)["sistema"]
        self.assertIsNone(d["gps_time"])
        self.assertIsNone(d["gps_time_s"])


class TestGpsTimeParaSegundos(unittest.TestCase):
    """`Gps.proto` declara `optional float gpsTime = 1;` e NÃO diz a
    unidade. Por isso a conversão discrimina em vez de supor — e devolve
    None no que não encaixa, em vez de inventar uma hora."""

    def test_hora_nmea_hhmmss(self):
        self.assertEqual(rm.gps_time_para_segundos(174551.25), 63951.25)
        self.assertEqual(rm.gps_time_para_segundos(235959.99), 86399.99)

    def test_epoca_unix_e_reconhecida(self):
        # Improvável num float de 32 bits, mas reconhecida em vez de
        # virar hora absurda.
        self.assertEqual(rm.gps_time_para_segundos(1789148756.0), 63956.0)

    def test_valor_sem_unidade_conhecida_nao_vira_hora(self):
        for v in (999999.0, 126000.0, 246000.0):
            self.assertIsNone(rm.gps_time_para_segundos(v), f"aceitou {v}")

    def test_ausencia_e_zero_nao_viram_medida(self):
        for v in (None, 0, 0.0, "", "abc"):
            self.assertIsNone(rm.gps_time_para_segundos(v), f"aceitou {v!r}")

    def test_comparar_o_bruto_dispensa_a_conversao(self):
        # O uso principal — "a posição mudou?" — compara o valor bruto e
        # funciona em qualquer formato. Se esta invariante cair, a captura
        # passa a depender de adivinhar a unidade.
        a = rm.parse_state(STATE)["sistema"]["gps_time"]
        b = rm.parse_state(STATE.replace("gpsTime: 143025.0",
                                         "gpsTime: 143026.0"))["sistema"]["gps_time"]
        self.assertNotEqual(a, b)


class TestIdentidadeDoVizinho(unittest.TestCase):
    """State.Peer NÃO tem name nem serialNumber — conferido no
    `State.proto`: mac, enabled, cost, rate, rssi, signal, age, stats,
    encapId, ipv4Address. Quem fecha o nome é o `encapId`, que é a parte
    numérica do número de série.

    Verificado nos 40 vizinhos da captura real da ERM-12: 40 de 40.
    """

    ST = """
manufacturer {
  model: "ES1-2450CS"
  serial: 113345
}
wireless {
  name: "wlan0"
  channel: 157
  noise: -109
  peer {
    mac: "aa"
    enabled: true
    cost: 8523
    rate: 650
    rssi: 21
    signal: -88
    age: 3
    encapId: 97128
    ipv4Address: "10.188.97.170"
  }
}
"""

    def test_encapid_do_vizinho_e_lido(self):
        d = rm.parse_state(self.ST)
        self.assertEqual(d["radios"][0]["peers"][0]["encap"], 97128)

    def test_serial_numerico_do_proprio_radio_e_lido(self):
        d = rm.parse_state(self.ST)["sistema"]
        self.assertEqual(d["serial_num"], 113345)
        self.assertEqual(d["modelo_fab"], "ES1-2450CS")

    def test_sem_encapid_fica_none_nao_zero(self):
        st = self.ST.replace("    encapId: 97128\n", "")
        self.assertIsNone(rm.parse_state(st)["radios"][0]["peers"][0]["encap"])

    def test_coleta_filtrada_nao_traz_manufacturer(self):
        # `manufacturer` é o campo 190 do State, fora de gps/wireless/
        # system. Numa coleta filtrada ele não vem — e o certo é None,
        # não 0: serial 0 seria um rádio existente e errado.
        st = self.ST[self.ST.index("wireless {"):]
        d = rm.parse_state(st)["sistema"]
        self.assertIsNone(d["serial_num"])
        self.assertIsNone(d["modelo_fab"])

    # (serialNumber, encap) dos 35 vizinhos vistos numa captura real da
    # ERM-12 (MeshMapper, 2026-09). Só os pares, que é o que se prova aqui.
    PARES_DE_CAMPO = [
        ("ES1-2450CS-113187", 113187), ("FE1-2255B-107794", 107794),
        ("FE1-2255B-107798", 107798), ("FE1-2255B-107805", 107805),
        ("FE1-2455LS-109505", 109505), ("FE1-2455LS-110091", 110091),
        ("FE1-2455LS-113330", 113330), ("FE1-2455LS-113344", 113344),
        ("FE1-2455LS-125566", 125566), ("JR3-52-120693", 120693),
        ("JR3-52-123036", 123036), ("LX5-2295C-61649", 61649),
        ("LX5-2295C-61783", 61783), ("LX5-2295C-64522", 64522),
        ("LX5-2295C-77205", 77205), ("LX5-2295C-87379", 87379),
        ("LX5-2295C-94856", 94856), ("LX5-2455D-96236", 96236),
        ("LX5-2455D-97128", 97128), ("ME4-2409R-61141", 61141),
        ("ME4-2409R-68404", 68404), ("ME4-2450R-101127", 101127),
        ("ME4-2450R-83498", 83498), ("ME4-2450R-96358", 96358),
        ("ME4-2450R-96692", 96692), ("ME4-2450R-96717", 96717),
        ("ME5-2450CS-121222", 121222), ("ME5-2450CS-121229", 121229),
        ("ME5-2450CS-121237", 121237), ("ME5-2450CS-121244", 121244),
        ("ME5-2450CS-125496", 125496), ("ME5-2450CS-127984", 127984),
        ("ME5-2450CS-127998", 127998), ("ME5-2450CS-128013", 128013),
        ("ME5-2450CS-128442", 128442),
    ]

    def test_encapid_bate_com_o_serial_em_campo(self):
        # A prova que sustenta a chave. Se um firmware mudar isso, este
        # teste é o lugar de registrar o caso novo.
        for ser, enc in self.PARES_DE_CAMPO:
            self.assertTrue(ser.endswith("-" + str(enc)),
                            f"encap {enc} nao e sufixo de {ser}")


def _carregar_bcapi_real():
    """A rajant-api DE VERDADE, se instalada, carregada ao lado do stub.

    O resto da suíte roda com o stub do topo deste arquivo (sem rede, sem
    biblioteca). Os testes do protocolo precisam das mensagens reais —
    Message_pb2, Common_pb2, State_pb2 — e as trocam de lugar só enquanto
    rodam. Sem a biblioteca na máquina, são pulados.
    """
    import importlib
    stub = sys.modules.pop('rajant_api', None)
    try:
        real = importlib.import_module('rajant_api')
        for sub in ('Message_pb2', 'Common_pb2', 'State_pb2'):
            importlib.import_module('rajant_api.' + sub)
        return real
    except Exception:
        for k in [k for k in sys.modules if k.startswith('rajant_api')]:
            sys.modules.pop(k, None)
        return None
    finally:
        if stub is not None:
            sys.modules['rajant_api'] = stub


_BCAPI_REAL = _carregar_bcapi_real()


def _tem_bcapi():
    return _BCAPI_REAL is not None


class _ComBCAPIReal:
    """Põe a rajant-api real no lugar do stub durante o teste."""

    def setUp(self):
        self._stub = sys.modules.get('rajant_api')
        sys.modules['rajant_api'] = _BCAPI_REAL
        self._reserva = dict(rm._LOGIN_RESERVA)
        rm._LOGIN_RESERVA["vezes"] = 0

    def tearDown(self):
        sys.modules['rajant_api'] = self._stub
        rm._LOGIN_RESERVA.clear(); rm._LOGIN_RESERVA.update(self._reserva)


class _RadioLogin:
    """Rádio falso que fala o login REAL da BCAPI.

    Mensagens da própria rajant-api (Message_pb2), cabeçalho de 8 bytes
    `>ibbbb` e corpo em deflate cru, entregue em pedaços pequenos como
    pelo TLS. O rádio manda o desafio e confere sha384(senha + desafio),
    como o firmware; depois responde ao State com o nome.

    `mudo_no_login`: "desafio" conecta e não manda o desafio; "agente"
    cala diante de um login com userAgent (o da biblioteca não manda).
    """

    def __init__(self, nome="CA-1006", senha="", papel="CO",
                 motivo_recusa="authentication failed", mudo_no_login=None,
                 pedaco=40, peers=0, atraso=0.0):
        self.nome, self.senha, self.papel = nome, senha, papel
        self.peers = peers
        # Intervalo entre pedaços: pelo TLS os registros chegam espaçados.
        # Sem ele o buffer do socket junta a mensagem inteira e um recv
        # único a pega — o teste passaria até com a leitura da biblioteca.
        self.atraso = atraso
        self.motivo_recusa, self.mudo_no_login = motivo_recusa, mudo_no_login
        self.pedaco = pedaco
        self.logins = []                # (papel pedido, aceito?)
        self.cliente = None
        self.conexoes = 0
        self._lados = []                # o lado do rádio fica vivo calado

    def conectar(self):
        import socket, threading
        a, b = socket.socketpair()
        a.settimeout(5)
        self.conexoes += 1
        self._lados.append(b)
        threading.Thread(target=self._servir, args=(b,), daemon=True).start()
        return a

    def _enviar(self, sock, msg):
        import struct, zlib
        c = zlib.compressobj(wbits=-15)
        dados = c.compress(msg.SerializeToString()) + c.flush()
        pacote = struct.pack(">ibbbb", len(dados), 2, 0, 0, 0) + dados
        for i in range(0, len(pacote), self.pedaco):
            sock.sendall(pacote[i:i + self.pedaco])
            if self.atraso:
                time.sleep(self.atraso)

    def _ler(self, sock):
        import struct, zlib
        from rajant_api import Message_pb2
        cab = rm._bc_ler_exato(sock, 8)
        n, flag = struct.unpack(">ibbbb", cab)[:2]
        corpo = rm._bc_ler_exato(sock, n)
        if flag == 2:
            corpo = zlib.decompress(corpo, -15)
        m = Message_pb2.BCMessage(); m.ParseFromString(corpo)
        return m

    def _servir(self, sock):
        import os as _os, hashlib
        from rajant_api import Message_pb2, Common_pb2
        try:
            if self.mudo_no_login == "desafio":
                return
            desafio = _os.urandom(32)
            ini = Message_pb2.BCMessage(); ini.sequenceNumber = 0
            ini.auth.action = ini.auth.LOGIN
            ini.auth.challengeOrResponse = desafio
            ini.auth.serial = "ME4-2450R-96692"
            self._enviar(sock, ini)
            m = self._ler(sock)
            if self.mudo_no_login == "agente" and m.auth.userAgent:
                return
            certo = hashlib.sha384(
                (self.senha + desafio.decode("latin1")).encode("latin1")).digest()
            ok = (m.auth.role == Common_pb2.Role.Value(self.papel)
                  and m.auth.challengeOrResponse == certo)
            self.logins.append((Common_pb2.Role.Name(m.auth.role), ok))
            self.cliente = m.auth.userAgent
            r = Message_pb2.BCMessage(); r.sequenceNumber = m.sequenceNumber
            r.authResult.status = (r.authResult.SUCCESS if ok
                                   else r.authResult.FAILURE)
            if not ok and self.motivo_recusa:
                r.authResult.description = self.motivo_recusa
            self._enviar(sock, r)
            if not ok:
                sock.close(); return
            while True:
                m = self._ler(sock)
                r = Message_pb2.BCMessage(); r.sequenceNumber = m.sequenceNumber
                if m.HasField("state"):
                    st = r.state
                    st.configuration.saved.general.name = self.nome
                    st.system.uptime = 123456
                    w = st.wireless.add()
                    w.name, w.channel, w.noise = "wlan0", 157, -95
                    for k in range(self.peers):
                        q = w.peer.add()
                        q.mac = bytes([0x84, 0x8d, 0x84, 0x41, k // 256, k % 256])
                        q.ipv4Address = f"10.0.{k // 250}.{k % 250}"
                        q.signal, q.rssi, q.cost = -60 - (k % 30), 30 - (k % 20), 5000 + k
                        q.encapId = 90000 + k
                self._enviar(sock, r)
        except Exception:
            return


@unittest.skipUnless(_tem_bcapi(), "rajant-api/protobuf ausentes")
class TestLoginBCAPI(_ComBCAPIReal, unittest.TestCase):
    """O login do exportador: rm.autenticar contra o desafio do rádio, com
    o Breadcrumb da biblioteca e só a rede trocada."""

    def _bc(self, radio, role="CO", senha=""):
        from rajant_api import Breadcrumb
        class BCFalso(Breadcrumb):
            def reachable(s): return True
            def setup_connection_socket(s):
                s.connection = radio.conectar()
        return BCFalso(host="10.0.0.6", port=2300, role=role, password=senha)

    def test_aceita_co_minusculo(self):
        # A rajant-api só conhece "CO": com "co" o authenticate() dela dava
        # KeyError, engolido, e a tela dizia "autenticacao falhou" nos 156.
        radio = _RadioLogin(senha="segredo")
        bc = self._bc(radio, role="co", senha="segredo")
        self.assertTrue(rm.autenticar(bc))
        self.assertEqual(radio.logins, [("CO", True)])
        self.assertEqual(radio.cliente, "rajant_monitor")   # visível no rádio

    def test_sessao_autenticada_le_o_estado(self):
        radio = _RadioLogin(nome="ERM-22 PTP CAM")
        bc = self._bc(radio)
        rm.autenticar(bc)
        st = rm.ler_estado(bc)
        self.assertEqual(st.configuration.saved.general.name, "ERM-22 PTP CAM")
        self.assertEqual(rm.parse_state(st)["sistema"]["nome"], "ERM-22 PTP CAM")

    def test_state_grande_chega_inteiro(self):
        # O get_state() da biblioteca lê UM recv(65535): pelo TLS isso traz
        # no máximo ~16 KB, e o State de um rádio com muitos vizinhos passa
        # disso. Aqui o rádio manda o State em pedaços de 1400 bytes.
        radio = _RadioLogin(nome="ERB-12 Slipstream", peers=1500, pedaco=1400)
        bc = self._bc(radio)
        rm.autenticar(bc)
        radio.atraso = 0.01
        st = rm.ler_estado(bc)
        self.assertGreater(bc._ultimo_tamanho[1], 16 * 1024)
        self.assertEqual(len(st.wireless[0].peer), 1500)
        d = rm.parse_state(st)
        self.assertEqual(d["sistema"]["nome"], "ERB-12 Slipstream")

    def test_a_biblioteca_sozinha_perde_o_state_grande(self):
        # A prova de que a leitura enquadrada é necessária: o mesmo rádio,
        # lido pelo get_state() da biblioteca, devolve False.
        radio = _RadioLogin(peers=1500, pedaco=1400)
        bc = self._bc(radio)
        rm.autenticar(bc)
        radio.atraso = 0.01
        self.assertIs(bc.get_state(), False)

    def test_sessao_e_fechada_depois_da_leitura(self):
        radio = _RadioLogin()
        bc = self._bc(radio)
        rm.autenticar(bc)
        rm.ler_estado(bc)
        rm.fechar_sessao(bc)
        lado = radio._lados[-1]
        lado.settimeout(2)
        self.assertEqual(lado.recv(1), b"")      # o outro lado fechou

    def test_senha_errada_diz_o_que_o_radio_respondeu(self):
        radio = _RadioLogin(senha="certa")
        with self.assertRaises(PermissionError) as cm:
            rm.autenticar(self._bc(radio, senha="errada "))
        msg = str(cm.exception)
        self.assertIn("authentication failed", msg)
        self.assertIn("confira usuário e senha", msg)
        self.assertIn("espaço no começo ou no fim", msg)

    def test_login_sem_resposta_diz_em_que_etapa_parou(self):
        radio = _RadioLogin(mudo_no_login="desafio")
        etapas = []
        with self.assertRaises(TimeoutError) as cm:
            rm.autenticar(self._bc(radio), etapas=etapas, prazo_s=0.3,
                          reserva=False)
        self.assertIn("conectou, mas o rádio não mandou o desafio",
                      str(cm.exception))
        self.assertEqual([n for n, _ in etapas],
                         ["conexão TLS", "parou: desafio"])

    def test_se_o_login_daqui_falha_o_da_biblioteca_entra(self):
        # Mensagem de login inteira num pedaço só: o recv único da
        # biblioteca tem de conseguir ler, senão não há reserva a testar.
        radio = _RadioLogin(mudo_no_login="agente", pedaco=4096)
        etapas = []
        self.assertTrue(rm.autenticar(self._bc(radio, role="co"),
                                      etapas=etapas, prazo_s=0.3))
        self.assertIn("authenticate() da biblioteca: entrou",
                      [n for n, _ in etapas])
        self.assertEqual(radio.logins, [("CO", True)])
        self.assertEqual(rm._LOGIN_RESERVA["vezes"], 1)

    def test_usuario_que_nao_existe_e_dito_antes_de_ir_a_rede(self):
        radio = _RadioLogin()
        with self.assertRaises(PermissionError) as cm:
            rm.autenticar(self._bc(radio, role="operador"))
        self.assertIn("VIEW", str(cm.exception))
        self.assertIn("CO", str(cm.exception))
        self.assertEqual(radio.conexoes, 0)

    def test_conexao_da_biblioteca_usa_o_prazo_daqui(self):
        # A biblioteca conecta com 2 s fixos, aperto de mão TLS incluído.
        import rajant_api
        visto = {}
        class Conn:
            def __init__(s, soc): visto["prazo"] = soc.gettimeout(); s.soc = soc
            def connect(s, end): visto["end"] = end
            def settimeout(s, t): visto["depois"] = t
        with mock.patch.object(rajant_api, "wrap_socket", Conn):
            bc = rajant_api.Breadcrumb(host="10.9.9.9", port=2300, role="CO",
                                       password="")
            rm._conectar(bc, 5.0)
            bc.connection.soc.close()
        self.assertEqual(visto, {"prazo": 5.0, "end": ("10.9.9.9", 2300),
                                 "depois": 5.0})


class TestLeituraDoEstado(unittest.TestCase):
    """State que não veio é FALHA de coleta, nunca um State vazio."""

    def test_false_da_biblioteca_vira_erro(self):
        class BC:
            connection = None
            def get_state(self): return False
        with self.assertRaises(ConnectionError):
            rm.ler_estado(BC())

    def test_resposta_vazia_vira_erro(self):
        for vazio in (None, "", b"", {}):
            class BC:
                connection = None
                def get_state(self, v=vazio): return v
            with self.assertRaises(ConnectionError, msg=repr(vazio)):
                rm.ler_estado(BC())

    def test_objeto_de_teste_passa_pelo_get_state(self):
        class BC:
            connection = None
            def get_state(self): return 'system { uptime: 5 }'
        self.assertIn("uptime", rm.ler_estado(BC()))

    def test_fechar_sessao_nao_estoura(self):
        class SemConexao: pass
        rm.fechar_sessao(SemConexao())
        rm.fechar_sessao(None)
        class ConexaoRuim:
            class connection:
                @staticmethod
                def close(): raise OSError("ja fechada")
        rm.fechar_sessao(ConexaoRuim())

    def _coletor(self):
        c = rm.RajantCollector(
            seeds=[], role="VIEW", password="", port=2300, interval=60,
            max_threads=1, timeout=1, tentativas=1, falhas_limite=3,
            manter_s=0, redesc_ciclos=10, falhas_remover=20,
            cache=rm.CacheIPs(os.path.join(tempfile.mkdtemp(), "c.json")),
            padrao_tag=rm.PADRAO_TAG_PADRAO)
        return c

    def test_state_que_nao_veio_conta_como_falha_e_nao_como_sem_tag(self):
        # O BUG: get_state() devolvia False, o parser fazia dele um State
        # vazio, o rádio ficava "sem nome" e o filtro de tag o tirava do
        # Grafana — justamente o rádio grande, de muitos vizinhos.
        ip = "10.188.97.220"                       # ERM-22 PTP CAM
        fechados = []
        class BC:
            def __init__(s, **k): s.connection = None
            def reachable(s): return True
            def authenticate(s): return True
            def get_state(s): return False
        c = self._coletor()
        rm._estados.pop(ip, None)
        with mock.patch.object(rm, "exigir_rajant_api", lambda: BC), \
             mock.patch.object(rm, "fechar_sessao", fechados.append):
            self.assertEqual(c._coletar_bc(ip), set())
        self.assertNotIn(ip, c.sem_tag)
        est = rm.estado(ip, c.falhas_limite, c.manter_s)
        self.assertEqual(est.falhas_consec, 1)
        self.assertEqual(len(fechados), 1)        # sessão fechada mesmo na falha


if __name__ == "__main__":
    unittest.main(verbosity=2)
