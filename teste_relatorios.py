#!/usr/bin/env python3
"""Testes da separação survey/semanal e da identidade Anglo.

Isto NÃO é o `teste_parser.py` original (346 testes): aquele se perdeu na
reciclagem do ambiente e não é reconstituível de memória. Aqui está o que
protege a parte que voltou a ser mexida — a separação dos dois relatórios,
o filtro que decide o que sai do semanal e a repintura na identidade.

Roda sem rádio e sem a lib `rajant_api`.
"""
import sys, types, unittest, inspect, tempfile
from pathlib import Path
from unittest import mock

# ── rajant_api não é instalável em CI; o módulo faz sys.exit(1) sem ela ──
if "rajant_api" not in sys.modules:
    _fake = types.ModuleType("rajant_api")
    class Breadcrumb:                     # stub: nenhum teste vai à rede
        def __init__(self, *a, **k): pass
    _fake.Breadcrumb = Breadcrumb
    sys.modules["rajant_api"] = _fake

sys.argv = ["x"]
import rajant_monitor as rm

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


def _deck_de_secoes(titulos_e_subtitulos):
    """Deck com um slide por (título, subtítulo), no leiaute do template:
    o botão ◂ MENU em 0,28" — ACIMA do título, em 0,32"."""
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    vazio = p.slide_masters[0].slide_layouts[6]
    for titulo, sub in titulos_e_subtitulos:
        s = p.slides.add_slide(vazio)
        bt = s.shapes.add_textbox(Inches(12.05), Inches(0.28),
                                  Inches(1.05), Inches(0.34))
        bt.text_frame.paragraphs[0].add_run().text = "◂ MENU"
        t = s.shapes.add_textbox(Inches(0.55), Inches(0.32),
                                 Inches(11.3), Inches(0.6))
        t.text_frame.paragraphs[0].add_run().text = titulo
        u = s.shapes.add_textbox(Inches(0.55), Inches(0.92),
                                 Inches(11.3), Inches(0.32))
        u.text_frame.paragraphs[0].add_run().text = sub
    return p


class TestTituloDoSlide(unittest.TestCase):
    def test_botao_de_menu_nao_vira_titulo(self):
        # O ◂ MENU fica em 0,28" e o título em 0,32": pegar "o mais alto"
        # sem filtrar botão devolveria o botão.
        p = _deck_de_secoes([("1. Dashboard Executivo", "resumo")])
        self.assertEqual(rm._titulo_do_slide(list(p.slides)[0]),
                         "1. Dashboard Executivo")

    def test_slide_sem_texto_nao_estoura(self):
        p = Presentation()
        s = p.slides.add_slide(p.slide_masters[0].slide_layouts[6])
        self.assertEqual(rm._titulo_do_slide(s), "")


class TestRemocaoDaSecaoDeSurvey(unittest.TestCase):
    """O semanal remove a seção de survey que vem no template. O que se
    testa aqui é que ele remove SÓ ela."""

    SECOES = [
        ("Navegação", "Clique numa seção para ir ao slide"),
        ("1. Dashboard Executivo", "Resumo da saúde da infraestrutura"),
        ("4. Saúde da Mesh Rajant — Links Críticos (SNR)",
         "10 piores enlaces do período — candidatos a realinhamento / "
         "site survey  |  gerado automaticamente"),
        ("4. Saúde da Mesh Rajant — Espectro por Canal",
         "Ocupação e ruído agregados por canal — todos os rádios ativos"),
        ("5. Site Survey — Intensidade de Sinal (RSSI) — 2.4 GHz", "medições"),
        ("5. Site Survey — Análise, Recomendações e Conclusão", "conclusão"),
        ("6. Manutenções", "preventivas e corretivas do período"),
    ]

    def _restantes(self):
        p = _deck_de_secoes(self.SECOES)
        n = rm.remover_slides_survey(p)
        return n, [rm._titulo_do_slide(s) for s in p.slides]

    def test_tira_os_slides_de_survey(self):
        n, restantes = self._restantes()
        self.assertEqual(n, 2)
        self.assertFalse([t for t in restantes if t.startswith("5. Site Survey")])

    def test_prosa_sobre_survey_no_subtitulo_nao_apaga_o_slide(self):
        # A REGRESSÃO: o filtro casava "site survey" em qualquer texto do
        # slide. O subtítulo de Links Críticos diz "candidatos a
        # realinhamento / site survey" e o slide sumia do semanal.
        _, restantes = self._restantes()
        self.assertIn("4. Saúde da Mesh Rajant — Links Críticos (SNR)",
                      restantes)

    def test_espectro_por_canal_fica_no_semanal(self):
        # A OUTRA METADE: o slide vinha titulado "5. Site Survey —
        # Espectro por Canal" e ia junto, mas os números são dos
        # contadores do exporter. Era o único gráfico de espectro do
        # semanal.
        _, restantes = self._restantes()
        self.assertIn("4. Saúde da Mesh Rajant — Espectro por Canal",
                      restantes)

    def test_navegacao_sobrevive(self):
        # O menu cita TODAS as seções e casaria com o filtro.
        _, restantes = self._restantes()
        self.assertIn("Navegação", restantes)

    def test_o_resto_do_relatorio_fica_de_pe(self):
        _, restantes = self._restantes()
        for t in ("1. Dashboard Executivo", "6. Manutenções"):
            self.assertIn(t, restantes)

    def test_deck_sem_survey_nao_perde_nada(self):
        p = _deck_de_secoes([("1. Dashboard Executivo", "resumo"),
                             ("6. Manutenções", "preventivas")])
        self.assertEqual(rm.remover_slides_survey(p), 0)
        self.assertEqual(len(list(p.slides)), 2)

    def test_botao_do_menu_deixa_de_apontar_para_o_vazio(self):
        fonte = inspect.getsource(rm.remover_slides_survey)
        self.assertIn("relatório separado", fonte)
        self.assertIn("_limpar_link_de_texto", fonte)


class TestSlidesTecnicosNoSemanal(unittest.TestCase):
    def test_espectro_pertence_a_secao_4(self):
        fonte = inspect.getsource(rm._slides_tecnicos)
        self.assertIn('"4. Saúde da Mesh Rajant — Espectro por Canal"', fonte)
        self.assertNotIn('"5. Site Survey — Espectro por Canal"', fonte)

    def test_espectro_ancora_em_slide_que_existe_no_semanal(self):
        # Ancorado só em slides de survey, ele ficava solto no fim do
        # deck quando a seção 5 não existia.
        fonte = inspect.getsource(rm._slides_tecnicos)
        trecho = fonte[fonte.index("Espectro por Canal"):]
        chamada = trecho[trecho.index("mover_apos_titulo("):]
        chamada = chamada[:chamada.index(")")]
        self.assertTrue("Ethernet e Quedas" in chamada
                        or "Links Críticos" in chamada,
                        f"ancora so em slide de survey: {chamada}")


class TestIdentidadeAnglo(unittest.TestCase):
    """A repintura do template do cliente na identidade Anglo."""

    def _deck(self):
        p = Presentation()
        p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
        vazio = p.slide_masters[0].slide_layouts[6]

        def texto(s, x, y, w, h, txt, cor, tam=12):
            tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            r = tb.text_frame.paragraphs[0].add_run(); r.text = txt
            r.font.name = "Calibri"; r.font.size = Pt(tam)
            r.font.color.rgb = RGBColor.from_string(cor)
            return tb

        capa = p.slides.add_slide(vazio)
        capa.background.fill.solid()
        capa.background.fill.fore_color.rgb = RGBColor.from_string("0D2052")
        texto(capa, 0.9, 2.3, 11.5, 0.9, "RELATÓRIO SEMANAL", "FFFFFF", 40)
        # Tom de corpo sobre a capa: vira cinza e some no azul se ninguém
        # olhar a luminância do fundo NOVO.
        texto(capa, 0.9, 3.25, 11.5, 0.45,
              "Infraestrutura de Telecom | Tecnologia de Mina", "A0B0CC", 16)

        s = p.slides.add_slide(vazio)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string("0D2052")
        texto(s, 0.55, 0.32, 11.3, 0.6, "1. Dashboard Executivo", "FFFFFF", 28)
        texto(s, 0.55, 0.92, 11.3, 0.32, "Resumo da semana", "A0B0CC")
        cart = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55),
                                  Inches(1.55), Inches(4.0), Inches(1.5))
        cart.fill.solid()
        cart.fill.fore_color.rgb = RGBColor.from_string("1A2744")
        cart.line.color.rgb = RGBColor.from_string("2A3A5C")
        # Rótulo BRANCO sobre cartão que vai clarear: tem de virar escuro.
        texto(s, 0.75, 1.70, 3.6, 0.3, "DISPONIBILIDADE DA MALHA", "FFFFFF", 11)
        texto(s, 0.75, 2.05, 3.6, 0.85, "99,12%", "C8D0E0", 30)
        gf = s.shapes.add_table(2, 2, Inches(0.55), Inches(3.3),
                                Inches(6.0), Inches(0.8))
        for j, cab in enumerate(("Indicador", "Valor")):
            cel = gf.table.cell(0, j)
            cel.fill.solid()
            cel.fill.fore_color.rgb = RGBColor.from_string("1A3A7A")
            r = cel.text_frame.paragraphs[0].add_run(); r.text = cab
            r.font.color.rgb = RGBColor.from_string("FFFFFF")
        for j, v in enumerate(("Disponibilidade", "OK")):
            cel = gf.table.cell(1, j)
            cel.fill.solid()
            cel.fill.fore_color.rgb = RGBColor.from_string("152238")
            r = cel.text_frame.paragraphs[0].add_run(); r.text = v
            r.font.color.rgb = RGBColor.from_string("27AE60")
        bt = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(12.05),
                                Inches(0.28), Inches(1.05), Inches(0.34))
        bt.fill.solid(); bt.fill.fore_color.rgb = RGBColor.from_string("1A2744")
        r = bt.text_frame.paragraphs[0].add_run(); r.text = "◂ MENU"
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
        bt.click_action.target_slide = capa
        return p

    def test_fundo_de_conteudo_vira_branco_e_a_capa_vira_azul(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        sl = list(p.slides)
        self.assertEqual(str(sl[0].background.fill.fore_color.rgb),
                         rm.ANGLO["azul"])
        self.assertEqual(str(sl[1].background.fill.fore_color.rgb),
                         rm.ANGLO["fundo"])

    def test_nao_sobra_cor_estrutural_do_template_escuro(self):
        import io, re, zipfile, collections
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        buf = io.BytesIO(); p.save(buf)
        z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
        achadas = collections.Counter()
        for n in z.namelist():
            if not re.match(r"ppt/(slides|charts)/.*\.xml$", n): continue
            for c in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"',
                                z.read(n).decode("utf8")):
                achadas[c.upper()] += 1
        sobrou = {c: k for c, k in achadas.items() if c in rm.ANGLO_FUNDOS}
        self.assertEqual(sobrou, {}, f"cor do template escuro sobrou: {sobrou}")

    def test_texto_claro_vira_escuro_sobre_o_fundo_branco(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        for sh in list(p.slides)[1].shapes:
            if sh.has_text_frame and "99,12" in sh.text_frame.text:
                self.assertEqual(
                    str(sh.text_frame.paragraphs[0].runs[0].font.color.rgb),
                    rm.ANGLO["texto"])
                return
        self.fail("o valor do cartão sumiu do slide")

    def test_branco_continua_branco_onde_o_fundo_seguiu_escuro(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        for sh in list(p.slides)[0].shapes:
            if sh.has_text_frame and "RELATÓRIO" in sh.text_frame.text:
                self.assertEqual(
                    str(sh.text_frame.paragraphs[0].runs[0].font.color.rgb),
                    "FFFFFF")
                return
        self.fail("título da capa sumiu")

    def test_borda_de_celula_e_repintada(self):
        # As bordas moram em lnL/lnR/lnT/lnB dentro do tcPr e NÃO passam
        # pelo python-pptx: sem tratá-las sobra a grade azul-escura
        # riscando o fundo branco.
        from lxml import etree
        p = self._deck()
        tab = [sh for sh in list(p.slides)[1].shapes if sh.has_table][0].table
        tc = tab.cell(0, 0)._tc
        tcPr = tc.find(f"{rm._A_NS}tcPr")
        if tcPr is None:
            tcPr = etree.SubElement(tc, f"{rm._A_NS}tcPr")
        ln = etree.SubElement(tcPr, f"{rm._A_NS}lnL")
        sf = etree.SubElement(ln, f"{rm._A_NS}solidFill")
        etree.SubElement(sf, f"{rm._A_NS}srgbClr").set("val", "2A3A5C")
        rm.aplicar_identidade_anglo(p)
        cores = [c.get("val") for c in tab._tbl.iter(f"{rm._A_NS}srgbClr")]
        self.assertNotIn("2A3A5C", cores)
        self.assertIn(rm.ANGLO["linha"], cores)

    def test_o_link_do_botao_sobrevive(self):
        p = self._deck(); rm.aplicar_identidade_anglo(p)
        ligados = 0
        for sh in list(p.slides)[1].shapes:
            try:
                if sh.click_action.target_slide is not None: ligados += 1
            except Exception:
                pass
        self.assertEqual(ligados, 1, "a repintura levou o hyperlink junto")

    def test_nenhum_texto_se_perde(self):
        p = self._deck()
        antes = sorted(sh.text_frame.text for s in p.slides
                       for sh in s.shapes if sh.has_text_frame)
        rm.aplicar_identidade_anglo(p)
        depois = sorted(sh.text_frame.text for s in p.slides
                        for sh in s.shapes if sh.has_text_frame)
        for t in antes:
            self.assertIn(t, depois)

    def test_regua_nao_entra_por_cima_de_conteudo(self):
        p = self._deck()
        s = list(p.slides)[1]
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55),
                           Inches(rm.ANGLO_SEMANAL["regua_y"] - 0.05),
                           Inches(5.0), Inches(0.3))
        rm.aplicar_identidade_anglo(p)
        self.assertFalse(any(sh.name == "anglo_regua" for sh in s.shapes),
                         "régua desenhada em cima do conteúdo")

    def test_slide_que_ja_nasceu_anglo_nao_ganha_segunda_marca(self):
        p = self._deck()
        n1 = rm.aplicar_identidade_anglo(p)["marcas"]
        n2 = rm.aplicar_identidade_anglo(p)["marcas"]
        self.assertGreater(n1, 0)
        self.assertEqual(n2, 0, "a segunda passagem marcou de novo")

    def test_todo_texto_fica_legivel_sobre_o_que_ficou_atras(self):
        def lum(h):
            def canal(i):
                v = int(h[i:i+2], 16) / 255
                return v/12.92 if v <= 0.03928 else ((v+0.055)/1.055)**2.4
            return 0.2126*canal(0) + 0.7152*canal(2) + 0.0722*canal(4)
        def razao(a, b):
            la, lb = lum(a), lum(b)
            return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)

        p = self._deck(); rm.aplicar_identidade_anglo(p)
        piores = []
        for s in p.slides:
            formas = list(rm._formas_planas(s.shapes))
            fundo_s = str(s.background.fill.fore_color.rgb).upper()
            alvos = []
            for sh in formas:
                if sh.has_text_frame:
                    f = rm._hex_do_fill(sh)
                    if f is None:
                        atras = rm._cartao_atras(sh, formas)
                        f = rm._hex_do_fill(atras) if atras is not None else fundo_s
                    alvos.append((sh.text_frame, f))
                if sh.has_table:
                    for row in sh.table.rows:
                        for cel in row.cells:
                            alvos.append((cel.text_frame,
                                          rm._hex_do_fill(cel) or fundo_s))
            for tf, fd in alvos:
                for par in tf.paragraphs:
                    for r in par.runs:
                        try: hx = str(r.font.color.rgb).upper()
                        except Exception: continue
                        rz = razao(hx, fd or "FFFFFF")
                        if rz < 3.0:
                            piores.append((round(rz, 2), hx, fd, r.text[:20]))
        self.assertEqual(piores, [],
                         f"texto ilegível após a repintura: {piores}")


class TestLigacaoComAGeracao(unittest.TestCase):
    def test_identidade_aplicada_por_padrao_e_por_ultimo(self):
        fonte = inspect.getsource(rm.gerar_ppt)
        self.assertIn('"identidade_anglo", fallback=True', fonte)
        # Repintar antes de preencher deixaria de fora o verde/âmbar/
        # vermelho do status e os slides técnicos.
        self.assertLess(fonte.index("_slides_tecnicos(p, d, cfg"),
                        fonte.index("aplicar_identidade_anglo(p)"))
        self.assertLess(fonte.index("anexar_survey_ao_ppt(p, cfg, survey)"),
                        fonte.index("aplicar_identidade_anglo(p)"))

    def test_semanal_nao_anexa_survey_por_padrao(self):
        fonte = inspect.getsource(rm.gerar_ppt)
        self.assertIn('"survey_no_semanal", fallback=False', fonte)
        self.assertLess(fonte.index("remover_slides_survey(p)"),
                        fonte.index("anexar_survey_ao_ppt(p, cfg, survey)"))

    def test_chaves_de_config_existem(self):
        cfg = rm.configparser.ConfigParser()
        with mock.patch.object(rm, "CONFIG_FILE",
                               str(Path(tempfile.mkdtemp()) / "c.ini")):
            rm.cfg_relatorio(cfg)
        for k in ("identidade_anglo", "survey_no_semanal"):
            self.assertTrue(cfg.has_option("relatorio", k), k)


if __name__ == "__main__":
    unittest.main(verbosity=2)
